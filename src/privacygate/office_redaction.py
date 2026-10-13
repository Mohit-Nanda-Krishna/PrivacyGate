"""Redacted copies of DOCX and PPTX files built from pipeline results.

redact_office() returns a copy of the source file in which every value the
final sanitized text replaced is replaced by the same token, everywhere in the
package, and the places a reader does not see on screen are cleaned too:

* text: every paragraph of every XML part (body, tables, headers/footers, text
  boxes, footnotes, comments, slides, speaker notes, charts), including
  deleted tracked-change text and field codes. A value split over several
  formatted runs is replaced across them; untouched text keeps its formatting.
  The replacements are the (original -> token) pairs of the sanitized text,
  applied wherever that original appears, so a name detected once is also
  replaced in a comment the extractor never read;
* attributes: image alt text and titles, mailto/hyperlink targets, and the
  author/initials of comments and tracked changes;
* document properties: author, last-modified-by, title, subject, keywords,
  description, company, manager and all custom properties are cleared;
* images: with mask_all_images (strict, the default) every raster image is
  replaced by a black image of the same size, because a screenshot can show
  values OCR did not read. Otherwise only the OCR'd words that were replaced in
  the sanitized text are blacked out, and images whose OCR failed are masked
  whole. Vector images (EMF/WMF/SVG) cannot be OCR'd or masked and are counted.

After writing, the copy is read back and every replaced value is searched for
again; the count is reported as `leftover` and must be 0. The caller must not
treat a copy with leftovers, or with unmasked vector images, as fully redacted.
"""

from __future__ import annotations

import io
import posixpath
import re
import zipfile
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw

from privacygate.models import ContentBlock, Document

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
# Paragraph elements and the text nodes inside them, per namespace.
PARAGRAPHS = {f"{{{W}}}p": (f"{{{W}}}t", f"{{{W}}}delText", f"{{{W}}}instrText"),
              f"{{{A}}}p": (f"{{{A}}}t",)}
# Breaks and tabs separate words without being text nodes; they count as whitespace.
SEPARATORS = {f"{{{W}}}br": "\n", f"{{{W}}}cr": "\n", f"{{{W}}}tab": "\t", f"{{{A}}}br": "\n"}


def _paragraph_nodes(paragraph: etree._Element, text_tags: tuple[str, ...]) -> list[tuple[etree._Element | None, str]]:
    """(text node or None for a separator, its text) in document order."""
    nodes = []
    for node in paragraph.iter(*text_tags, *SEPARATORS):
        if node.tag in SEPARATORS:
            # A w:tab inside w:tabs is a tab stop definition, not text.
            if etree.QName(node.getparent()).localname != "tabs":
                nodes.append((None, SEPARATORS[node.tag]))
        else:
            nodes.append((node, node.text or ""))
    return nodes
TOKEN = re.compile(r"\[[A-Z][A-Z_]*_\d+\]")
MIN_VALUE = 3  # letters/digits a replaced value needs before it is swept through the package
PERSON_ATTRIBUTES = frozenset({"author", "initials", "userId", "providerId"})
SWEPT_ATTRIBUTES = frozenset({"descr", "title", "name", "Target", "tooltip"})
VECTOR = frozenset({".emf", ".wmf", ".svg"})
CORE_KEPT = frozenset({"created", "modified", "revision"})


@dataclass
class RedactedOffice:
    data: bytes
    file_type: str
    replacements: int = 0  # text spans replaced in the package
    parts_changed: int = 0  # XML parts with at least one change
    attributes_cleaned: int = 0  # alt text, link targets, comment/revision authors
    images_masked_whole: int = 0
    image_words_masked: int = 0
    vector_images: int = 0  # vector images left as they are (cannot be OCR'd or masked)
    unmasked_images: int = 0  # raster images that could not be decoded, left as they are
    leftover: int = 0  # replaced values still found in the copy after writing (must be 0)
    mask_all_images: bool = True
    values: int = 0  # distinct replaced values swept through the package


def _parts(text: str) -> list[str]:
    return re.findall(r"\S+|\s+", text)


def _aligned_values(original: str, sanitized: str) -> list[tuple[str, str]] | None:
    """(value, token) for every token, by matching the text between tokens literally.

    The sanitized text is the original with each value replaced by a token, so the
    stretches between tokens appear unchanged and in order in the original.
    """
    tokens = TOKEN.findall(sanitized)
    if not tokens:
        return []
    pieces = TOKEN.split(sanitized)
    pattern = "(.+?)".join(re.escape(piece) for piece in pieces)
    match = re.fullmatch(pattern, original, re.S)
    if not match:
        return None
    return [(value.strip(), token) for value, token in zip(match.groups(), tokens)]


def replacement_pairs(document: Document, sanitized: Document) -> dict[str, str]:
    """original value -> token text for every value the sanitized text replaced."""
    final = {block.block_id: block.text for block in sanitized.blocks}
    pairs: dict[str, str] = {}
    for block in document.blocks:
        new = final.get(block.block_id)
        if new is None or new == block.text:
            continue
        aligned = _aligned_values(block.text, new)
        if aligned is not None:
            for value, token in aligned:
                if sum(character.isalnum() for character in value) >= MIN_VALUE:
                    pairs.setdefault(value, token)
            continue
        # Fallback when the literal alignment fails: word-level diff.
        old_parts, new_parts = _parts(block.text), _parts(new)
        matcher = SequenceMatcher(None, old_parts, new_parts, autojunk=False)
        for tag, i1, i2, j1, j2 in matcher.get_opcodes():
            if tag != "replace":
                continue
            old, token = "".join(old_parts[i1:i2]).strip(), "".join(new_parts[j1:j2]).strip()
            if not TOKEN.search(token):
                continue
            # Drop punctuation both sides share ("Lindqvist," -> "[PERSON_004],").
            while old and token and old[-1] == token[-1] and not old[-1].isalnum():
                old, token = old[:-1], token[:-1]
            while old and token and old[0] == token[0] and not old[0].isalnum():
                old, token = old[1:], token[1:]
            for value, replacement in _split_pair(old, token):
                if sum(character.isalnum() for character in value) >= MIN_VALUE:
                    pairs.setdefault(value, replacement)
    return pairs


def _split_pair(old: str, new: str) -> list[tuple[str, str]]:
    """One pair per token: "Wilson, Laura Mitchell" -> "[PERSON_1], [PERSON_2]" gives two pairs."""
    tokens = TOKEN.findall(new)
    if len(tokens) < 2:
        return [(old, new)]
    separators = TOKEN.split(new)
    pattern = "".join(re.escape(part) + ("(.+?)" if index < len(tokens) else "")
                      for index, part in enumerate(separators))
    match = re.fullmatch(pattern, old, re.S)
    if not match:
        return [(old, new)]
    return [(value.strip(), token) for value, token in zip(match.groups(), tokens)]


def _pattern(pairs: dict[str, str]) -> re.Pattern[str] | None:
    if not pairs:
        return None
    alternatives = sorted(pairs, key=len, reverse=True)
    return re.compile(r"(?<![\w])(?:" + "|".join(map(re.escape, alternatives)) + r")(?![\w])")


def _replace_runs(texts: list[str], pattern: re.Pattern[str], pairs: dict[str, str]) -> tuple[list[str], int]:
    """Replace values in a paragraph whose text is split over runs; each value's token goes in its first run."""
    matches = list(pattern.finditer("".join(texts)))
    if not matches:
        return texts, 0
    owner, run_start, offset = [], [], 0
    for index, text in enumerate(texts):
        run_start.append(offset)
        owner += [index] * len(text)
        offset += len(text)
    pieces = [list(text) for text in texts]
    for match in matches:
        for position in range(match.start(), match.end()):
            pieces[owner[position]][position - run_start[owner[position]]] = ""
        first = owner[match.start()]
        pieces[first][match.start() - run_start[first]] = pairs[match.group()]
    return ["".join(piece) for piece in pieces], len(matches)


def _clean_xml(data: bytes, pattern: re.Pattern[str] | None, pairs: dict[str, str], name: str,
               report: RedactedOffice) -> bytes:
    root = etree.fromstring(data)
    changed = 0
    if pattern is not None:
        for tag, text_tags in PARAGRAPHS.items():
            for paragraph in root.iter(tag):
                nodes = _paragraph_nodes(paragraph, text_tags)
                if not any(node is not None for node, _ in nodes):
                    continue
                new_texts, count = _replace_runs([text for _, text in nodes], pattern, pairs)
                if count:
                    for (node, _), text in zip(nodes, new_texts):
                        if node is None:
                            continue  # a break or tab: an element, not text; it stays
                        node.text = text
                        if text != text.strip():
                            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    report.replacements += count
                    changed += count
    for element in root.iter():
        for key, value in list(element.attrib.items()):
            local = etree.QName(key).localname
            if local in PERSON_ATTRIBUTES and value:
                element.set(key, "Redacted" if local != "initials" else "R")
                report.attributes_cleaned += 1
                changed += 1
            elif local in SWEPT_ATTRIBUTES and pattern is not None and pattern.search(value):
                element.set(key, pattern.sub(lambda m: pairs[m.group()], value))
                report.attributes_cleaned += 1
                changed += 1
    if not changed:
        return data
    report.parts_changed += 1
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _clean_rels(data: bytes, pattern: re.Pattern[str] | None, pairs: dict[str, str], report: RedactedOffice) -> bytes:
    """Hyperlink targets (mailto:...) live in relationship files."""
    if pattern is None:
        return data
    root = etree.fromstring(data)
    changed = False
    for element in root.iter():
        target = element.get("Target")
        if target and pattern.search(target):
            element.set("Target", pattern.sub(lambda m: pairs[m.group()], target))
            report.attributes_cleaned += 1
            changed = True
    if not changed:
        return data
    report.parts_changed += 1
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _clear_properties(name: str, data: bytes) -> bytes:
    root = etree.fromstring(data)
    if name == "docProps/custom.xml":
        for child in list(root):
            root.remove(child)
    else:
        # Remove rather than empty: every property is optional, but Word rejects some empty
        # values (an empty date such as cp:lastPrinted makes the file "corrupted").
        for element in list(root):
            local = etree.QName(element).localname
            if (name == "docProps/core.xml" and local not in CORE_KEPT) or (
                    name == "docProps/app.xml" and local in ("Company", "Manager")):
                root.remove(element)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _mask_image(data: bytes, extension: str, boxes: list[list[int]] | None) -> bytes:
    """Black out boxes (or the whole image when boxes is None), keeping size and format."""
    with Image.open(io.BytesIO(data)) as source:
        image_format = source.format or extension.lstrip(".").upper()
        image = source.convert("RGB")
    draw = ImageDraw.Draw(image)
    if boxes is None:
        draw.rectangle((0, 0, image.width, image.height), fill=(0, 0, 0))
    else:
        for left, top, right, bottom in boxes:
            draw.rectangle((left - 1, top - 1, right + 1, bottom + 1), fill=(0, 0, 0))
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG" if image_format.upper() in ("JPG", "JPEG") else image_format)
    return buffer.getvalue()


def _word_boxes(block: ContentBlock, sanitized_text: str) -> list[list[int]]:
    """Image-pixel boxes of the OCR words that the sanitized text replaced."""
    boxes = block.metadata.get("word_boxes_pixels") or []
    words = block.text.split(" ")
    if len(words) != len(boxes):
        return [box for box in boxes]  # cannot align: mask the whole line
    scale = block.metadata.get("ocr_scale") or 1
    old_parts, new_parts = _parts(block.text), _parts(sanitized_text)
    replaced: set[int] = set()
    matcher = SequenceMatcher(None, old_parts, new_parts, autojunk=False)
    word_index = {part_index: count for count, part_index in
                  enumerate(i for i, part in enumerate(old_parts) if not part.isspace())}
    for tag, i1, i2, _, _ in matcher.get_opcodes():
        if tag in ("replace", "delete"):
            replaced |= {word_index[i] for i in range(i1, i2) if i in word_index}
    return [[round(v / scale) for v in boxes[i]] for i in sorted(replaced)]


def _images(document: Document, sanitized: Document, mask_all: bool) -> dict[str, list[list[int]] | None]:
    """image name -> boxes to mask (None: the whole image)."""
    final = {block.block_id: block.text for block in sanitized.blocks}
    status = {entry.get("image"): entry.get("status") for entry in document.metadata.get("embedded_images", [])
              if isinstance(entry, dict)}
    plan: dict[str, list[list[int]] | None] = {name: None for name, state in status.items()
                                                if state == "failed" or mask_all}
    for block in document.blocks:
        name = block.metadata.get("image_name")
        if block.extraction_method != "embedded_image_ocr" or not name or plan.get(name, []) is None:
            continue
        if final.get(block.block_id, block.text) != block.text:
            plan.setdefault(name, []).extend(_word_boxes(block, final[block.block_id]))
    return plan


def package_text(data: bytes) -> str:
    """Every paragraph (runs joined, as a reader sees it) and attribute value in every XML part."""
    texts = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in archive.namelist():
            if not name.endswith((".xml", ".rels")):
                continue
            root = etree.fromstring(archive.read(name))
            for tag, text_tags in PARAGRAPHS.items():
                texts += ["".join(text for _, text in _paragraph_nodes(paragraph, text_tags))
                          for paragraph in root.iter(tag)]
            texts += [value for element in root.iter() for value in element.attrib.values()]
            texts += [element.text for element in root.iter() if isinstance(element.tag, str) and element.text]
    return "\n".join(texts)


def _add_watermark(data: bytes, file_type: str, watermark: str) -> bytes:
    buffer = io.BytesIO()
    if file_type == "docx":
        from docx import Document as WordDocument
        from docx.shared import RGBColor
        word = WordDocument(io.BytesIO(data))
        paragraph = word.paragraphs[0].insert_paragraph_before() if word.paragraphs else word.add_paragraph()
        run = paragraph.add_run(watermark)
        run.bold = True
        run.font.color.rgb = RGBColor(0xC0, 0, 0)
        word.save(buffer)
    else:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
        deck = Presentation(io.BytesIO(data))
        for slide in deck.slides:
            box = slide.shapes.add_textbox(Inches(0.2), Inches(0.05), deck.slide_width - Inches(0.4), Inches(0.4))
            run = box.text_frame.paragraphs[0].add_run()
            run.text = watermark
            run.font.size, run.font.bold, run.font.color.rgb = Pt(12), True, RGBColor(0xC0, 0, 0)
        deck.save(buffer)
    return buffer.getvalue()


def redact_office(
    source: str | Path, document: Document, sanitized: Document, watermark: str | None = None,
    mask_all_images: bool = True,
) -> RedactedOffice:
    """A redacted copy of a DOCX or PPTX; see the module docstring for what is cleaned."""
    file_type = Path(source).suffix.lower().lstrip(".")
    if file_type not in ("docx", "pptx"):
        raise ValueError("redact_office handles .docx and .pptx files.")
    pairs = replacement_pairs(document, sanitized)
    pattern = _pattern(pairs)
    report = RedactedOffice(b"", file_type, mask_all_images=mask_all_images, values=len(pairs))
    image_plan = _images(document, sanitized, mask_all_images)
    output = io.BytesIO()
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in archive.infolist():
            name, data = item.filename, archive.read(item.filename)
            base, extension = posixpath.basename(name), posixpath.splitext(name)[1].lower()
            if name in ("docProps/core.xml", "docProps/app.xml", "docProps/custom.xml"):
                data = _clear_properties(name, data)
            elif name.endswith(".rels"):
                data = _clean_rels(data, pattern, pairs, report)
            elif name.endswith(".xml") and name != "[Content_Types].xml":
                data = _clean_xml(data, pattern, pairs, name, report)
            elif "/media/" in name and extension in VECTOR:
                report.vector_images += 1
            elif "/media/" in name and (mask_all_images or base in image_plan):
                boxes = None if mask_all_images else image_plan.get(base)
                if boxes is None or boxes:
                    try:
                        data = _mask_image(data, extension, boxes)
                    except Exception:
                        report.unmasked_images += 1  # undecodable: left as it is, reported
                    else:
                        if boxes is None:
                            report.images_masked_whole += 1
                        else:
                            report.image_words_masked += len(boxes)
            target.writestr(item, data)
    data = output.getvalue()
    if watermark:
        data = _add_watermark(data, file_type, watermark)
    text = package_text(data)
    report.leftover = sum(1 for value in pairs if re.search(r"(?<![\w])" + re.escape(value) + r"(?![\w])", text))
    report.data = data
    return report
