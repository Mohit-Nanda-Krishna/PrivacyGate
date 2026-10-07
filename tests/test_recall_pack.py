"""Recall/precision pack and independent gate checks; synthetic strings only."""

import pytest
from docx import Document as DocxDocument

from privacygate.detection import detect_pii, detect_pii_with_registry, detect_regex, detect_custom
from privacygate.detection.names import DEFAULT_ALLOWLIST, build_name_registry, clean_ner_entity
from privacygate.models import ContentBlock, Document, PIIEntity
from privacygate.validation import validate_privacy
from privacygate.validation.residual_checks import independent_residual_checks


def _doc(*texts, **metadata_by_index):
    return Document("d1", "synthetic.pdf", "pdf", blocks=[
        ContentBlock(f"b{i}", text, page_number=1, metadata=metadata_by_index.get(f"m{i}", {}))
        for i, text in enumerate(texts)
    ])


def _found(document, entities):
    texts = {block.block_id: block.text for block in document.blocks}
    return {(e.entity_type, texts[e.block_id][e.start:e.end]) for e in entities}


def _values(entities, text):
    return {text[e.start:e.end] for e in entities}


# --- Name registry sweep -------------------------------------------------------

def test_name_sweep_finds_variants_of_a_detected_person():
    doc = _doc(
        "Report owner: Quentin R. Abernathy, Head of Widgets.",
        "Escalate to Q. Abernathy, or to Q. R. Abernathy out of hours.",
        "Abernathy approved it; Quentin signed later.",
        "Contact q.abernathy@example.test for queries.",
    )
    found = _found(doc, detect_pii(doc))
    for variant in ["Q. Abernathy", "Q. R. Abernathy", "Abernathy", "Quentin"]:
        assert ("PERSON", variant) in found
    sweeps = [e for e in detect_pii(doc) if "name_sweep" in e.detector]
    assert sweeps


def test_registry_from_email_local_part_and_honorific():
    blocks = [
        ContentBlock("b0", "Owner Zebedee Okafor-Lumen signs, see z.okaforlumen@example.test"),
        ContentBlock("b1", "Ms Thistlewood called about her account."),
    ]
    registry = build_name_registry(blocks, [])
    assert registry.find("Thanks, Ms Thistlewood") and registry.find("Okafor-Lumen agreed")


def test_registry_never_matches_inside_email_or_placeholder():
    registry = build_name_registry([ContentBlock("b", "Dear Mr Halvard")], [])
    assert registry.find("halvard@example.test") == []
    assert registry.find("[PERSON] replied") == []


# --- NER false positives ------------------------------------------------------

@pytest.mark.parametrize("text", [
    "Trigger Name: Asset Created", "Tier 2 Vendor", "Optiv", "Email", "WR.MP.6", "Change Workflow", "AI",
])
def test_ner_noise_is_dropped(text):
    block = ContentBlock("b", text)
    entity = PIIEntity("PERSON", 0, len(text), 0.85, "presidio:SpacyRecognizer", "b")
    assert clean_ner_entity(block, entity, DEFAULT_ALLOWLIST) == []


def test_ner_span_is_trimmed_and_split_at_line_breaks():
    text = "Tier Lucinda P. Marrow\nGideon Fairweather"
    block = ContentBlock("b", text)
    entity = PIIEntity("PERSON", 0, len(text), 0.85, "presidio:SpacyRecognizer", "b")
    cleaned = clean_ner_entity(block, entity, DEFAULT_ALLOWLIST)
    assert [text[e.start:e.end] for e in cleaned] == ["Lucinda P. Marrow", "Gideon Fairweather"]


def test_allowlist_is_configurable():
    doc = _doc("Prepared by Bramble Quayside for the Quayside programme.")
    assert ("PERSON", "Bramble Quayside") not in _found(doc, detect_pii(doc, allowlist={"Bramble"}))


# --- Missed structured patterns -------------------------------------------------

@pytest.mark.parametrize("text,value,entity_type", [
    ("M +1 (555) 0114", "+1 (555) 0114", "PHONE_NUMBER"),
    ("Mobile: 555 0199", "555 0199", "PHONE_NUMBER"),
    ("SSN 900-12-3456 on file", "900-12-3456", "US_SSN"),
    ("roster PAN ABCPD1234E listed", "ABCPD1234E", "GOVERNMENT_ID"),
    ("holder passport ZQ1234567 seen", "ZQ1234567", "GOVERNMENT_ID"),
    ("PESEL 900101xxxxx", "900101xxxxx", "GOVERNMENT_ID"),
    ("your card ending 4321 was charged", "4321", "ACCOUNT_NUMBER"),
    ("joint savings account ending in 9876.", "9876", "ACCOUNT_NUMBER"),
    ("last four digits of your social security number (1234)", "1234", "GOVERNMENT_ID"),
    ("DOB 01/02/1990 recorded", "01/02/1990", "DATE_OF_BIRTH"),
    ("date of birth, 4 March 1988.", "4 March 1988", "DATE_OF_BIRTH"),
    ("lives at 12 Maple Court, Apt 3B, Springfield, IL 62704 now", "12 Maple Court, Apt 3B, Springfield, IL 62704",
     "ADDRESS"),
    ("mail j.doe@exampletg.com or j.doe@example.com", None, None),
])
def test_structured_patterns(text, value, entity_type):
    entities = detect_regex(text)
    if value is None:
        return
    assert (entity_type, value) in {(e.entity_type, text[e.start:e.end]) for e in entities}


def test_card_and_label_words_are_not_redacted():
    text = "your card ending 4321"
    entities = [e for e in detect_regex(text) if e.entity_type == "ACCOUNT_NUMBER"]
    assert _values(entities, text) == {"4321"}


@pytest.mark.parametrize("text,value", [
    ("Owner EMP-12345 approved", "EMP-12345"),
    ("Badge EMP123456", "EMP123456"),
    ("Director DIR-0042", "DIR-0042"),
    ("OCR read D1R-0042 and DlR-0O42", "D1R-0042"),
    ("Vendor MER-IN-4471 and MER-PL-1187", "MER-IN-4471"),
    ("wrapped MER-\nPH-2214 id", "MER-\nPH-2214"),
    ("damaged EMF-3O104 badge", "EMF-3O104"),
])
def test_standalone_personnel_ids(text, value):
    assert value in _values(detect_custom(text), text)


def test_ocr_damaged_org_email_domain(monkeypatch):
    from privacygate.detection import regex_detector

    text = "contact smarcheaigfeadencetg.com today"
    monkeypatch.setattr(regex_detector, "OCR_EMAIL_PATTERN", regex_detector.re.compile(
        r"(?<![\w.@|])(?P<value>[A-Za-z0-9._|-]{2,40}\s?@?\s?" + regex_detector._ocr_domain("cadencefg.com")
        + r"m?)(?![\w])", regex_detector.re.IGNORECASE))
    assert "smarcheaigfeadencetg.com" in _values(detect_regex(text), text)


def test_table_header_gives_phone_context():
    doc = Document("d", "s.docx", "docx", blocks=[
        ContentBlock("h", "Mobile", metadata={"kind": "table_cell", "table_number": 1, "row_number": 1,
                                              "column_number": 2}),
        ContentBlock("c", "0412 555 019", metadata={"kind": "table_cell", "table_number": 1, "row_number": 2,
                                                    "column_number": 2}),
    ])
    assert ("PHONE_NUMBER", "0412 555 019") in _found(doc, detect_pii(doc))


def test_dates_and_references_are_not_pii():
    text = "Issue ISS-2026-0112 due 30 September 2026, see GRP-POL-001 section 14."
    doc = _doc(text)
    assert detect_pii(doc) == []


# --- Independent gate checks ----------------------------------------------------

def test_independent_checks_find_digits_emails_ids_and_registered_names():
    blocks = [ContentBlock("b", "Call 0207 946 0958, mail a@b.example, id EMP-77881 or ABCPD1234E; Q. Abernathy.")]
    registry = build_name_registry([ContentBlock("x", "Quentin R. Abernathy")], [("x", 0, 20)])
    checks = {e.detector for e in independent_residual_checks(blocks, registry)}
    assert checks == {"gate:digit_run", "gate:email_shape", "gate:id_shape", "gate:name_registry"}


def test_independent_checks_ignore_dates_references_and_placeholders():
    blocks = [ContentBlock("b", "On 2026-06-11 and 30/09/2026, ISS-2026-0112 for [PERSON] at [PHONE], page 12 of 35.")]
    assert independent_residual_checks(blocks, None) == []


def test_gate_blocks_on_independent_hit_even_if_detectors_are_blind(monkeypatch):
    monkeypatch.setattr("privacygate.validation.privacy_gate.detect_pii", lambda doc: [])
    doc = _doc("Reach the owner on 0207 946 0958.")
    result = validate_privacy(doc)
    assert result.status == "BLOCKED"
    assert "independent residual checks" in result.reason


def test_pii_missed_in_pass_one_is_never_approved_after_pass_two(tmp_path, monkeypatch):
    """Today's case: detectors miss PII, a second pass cleans what the gate listed, and
    the old logic approved. The independent checks must keep the document blocked."""
    import privacygate.pipeline as pipeline
    import privacygate.validation.privacy_gate as gate

    path = tmp_path / "case.docx"
    docx = DocxDocument()
    docx.add_paragraph("Owner Quentin R. Abernathy signed the record.")
    docx.add_paragraph("Escalate to Q. Abernathy on 0207 946 0958.")
    docx.save(path)

    def blind_detector(document, allowlist=None):
        # Finds only the full name: misses the variant and the phone number.
        entities = []
        for block in document.blocks:
            start = block.text.find("Quentin R. Abernathy")
            if start >= 0:
                entities.append(PIIEntity("PERSON", start, start + 20, 0.9, "presidio:test", block.block_id,
                                          {"paragraph_number": block.paragraph_number,
                                           "extraction_method": block.extraction_method}))
        return entities

    def blind_with_registry(document, allowlist=None):
        entities = blind_detector(document)
        texts = {b.block_id: b.text for b in document.blocks}
        return entities, build_name_registry(document.blocks, [(e.block_id, e.start, e.end) for e in entities])

    monkeypatch.setattr(pipeline, "detect_pii_with_registry", blind_with_registry)
    monkeypatch.setattr(gate, "detect_pii", blind_detector)

    result = pipeline.run_pipeline(path, max_passes=2)
    assert result.passes_executed == 2
    assert result.validation.status == "BLOCKED"
    assert "Independent residual checks" in result.validation.reason


def test_gate_still_approves_clean_output():
    doc = _doc("The [PERSON] approved the [EMAIL] change on 30 June 2026.")
    assert validate_privacy(doc).status == "APPROVED"


def test_detect_pii_with_registry_returns_registry():
    doc = _doc("Prepared by Ophelia Brandywine (o.brandywine@example.test).", "Brandywine reviewed it.")
    entities, registry = detect_pii_with_registry(doc)
    assert len(registry) == 1
    assert ("PERSON", "Brandywine") in _found(doc, entities)


def test_uncorroborated_name_pairs_are_not_swept():
    # Two capitalised words with no initial, e-mail, honorific or identifier nearby
    # (e.g. NER noise such as "Lagging Monthly") must not be swept across the document.
    blocks = [ContentBlock("b0", "Velvet Harbour"), ContentBlock("b1", "Harbour reviewed it.")]
    assert len(build_name_registry(blocks, [("b0", 0, 14)])) == 0
    assert len(build_name_registry(blocks, [("b0", 0, 14)], identifier_blocks=["b0"])) == 1


def _cell(block_id, text, row, column, table=1):
    return ContentBlock(block_id, text, metadata={"kind": "table_cell", "table_number": table,
                                                  "row_number": row, "column_number": column})


def test_role_header_column_names_including_two_letter_initials():
    doc = Document("d", "s.docx", "docx", blocks=[
        _cell("h1", "DESCRIPTION", 1, 1), _cell("h2", "USERS", 1, 2),
        _cell("c1", "Reviews workflows", 2, 1), _cell("c2", "Quilla Marsh\nXT Fenwick\nEdit", 2, 2),
        ContentBlock("img", "Created by Quilla Marsh 10/31/2025 Quilla Marsh 10/31/2025",
                     extraction_method="embedded_image_ocr", metadata={"image_name": "image1.png"}),
        ContentBlock("other", "Fenwick approved; Vendor Risk reviewed."),
    ])
    found = _found(doc, detect_pii(doc))
    assert {("PERSON", "Quilla Marsh"), ("PERSON", "XT Fenwick"), ("PERSON", "Fenwick")} <= found
    # Swept into image text even though NER never tagged it there.
    texts = {b.block_id: b.text for b in doc.blocks}
    image_hits = [e for e in detect_pii(doc) if e.block_id == "img" and e.entity_type == "PERSON"]
    assert len(image_hits) == 2
    assert ("PERSON", "Vendor Risk") not in found and ("PERSON", "Edit") not in found


def test_role_row_label_and_inline_label():
    doc = Document("d", "s.docx", "docx", blocks=[
        _cell("l", "Author(s)", 1, 1), _cell("v", "Ravindra Osterholt, Pell Grantham", 1, 2),
        ContentBlock("p", "Prepared by Ismay Corvell for the review."),
    ])
    found = _found(doc, detect_pii(doc))
    assert {("PERSON", "Ravindra Osterholt"), ("PERSON", "Pell Grantham"), ("PERSON", "Ismay Corvell")} <= found


def test_role_words_do_not_make_business_terms_names():
    doc = Document("d", "s.docx", "docx", blocks=[
        _cell("h", "Owner", 1, 1), _cell("c", "Vendor Risk Management Group", 2, 1),
    ])
    assert not [e for e in detect_pii(doc) if e.entity_type == "PERSON"]


def test_name_with_initial_shape_and_split_country_code():
    doc = Document("d", "s.pdf", "pdf", blocks=[
        ContentBlock("b0", "Roster: Ottoline V. Quarrington MER-IN-4471", page_number=1),
        ContentBlock("b1", "Deputy: on +1", page_number=1),
        ContentBlock("b2", "(917) 555-0177. Later.", page_number=1),
    ])
    found = _found(doc, detect_pii(doc))
    assert ("PERSON", "Ottoline V. Quarrington") in found
    assert ("PHONE_NUMBER", "+1") in found
