"""Native-only extraction contract, ordering, provenance, and failure behavior."""

from dataclasses import asdict
from io import BytesIO
from pathlib import Path
from shutil import copyfile
from uuid import UUID

import pymupdf
import pytest
from docx import Document as DocxDocument
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from privacygate.extraction import ExtractionError, OCRRequiredError, extract_document
from privacygate.extraction import docx, pdf, pptx
from privacygate.models import ContentBlock, Document

FIXTURES = Path(__file__).parent / "fixtures"
EXTRACTORS = {"pdf": pdf.extract_pdf, "docx": docx.extract_docx, "pptx": pptx.extract_pptx}


def test_pdf_text_order_and_page_sources():
    document = pdf.extract_pdf(FIXTURES / "native.pdf")
    assert [block.text for block in document.blocks] == [
        "Synthetic first page heading.", "Second block on first page.",
        "Synthetic second page heading.", "Final block on second page.",
    ]
    assert [block.page_number for block in document.blocks] == [1, 1, 2, 2]
    assert [block.metadata["block_order"] for block in document.blocks] == [1, 2, 1, 2]
    assert document.blocks[0].metadata["source_block_number"] == 1
    assert document.metadata["page_count"] == 2
    assert document.metadata["pages_without_native_text"] == []
    for block in document.blocks:
        assert block.extraction_method == "pdf_native"
        assert len(block.metadata["bbox"]) == 4
        assert block.slide_number is block.paragraph_number is None


@pytest.mark.parametrize("text", ["", "   ", "... --- !!!", "Short title", "a" * 19])
def test_pdf_with_insufficient_native_text_requires_attention(tmp_path, text):
    path = tmp_path / "sparse.pdf"
    with pymupdf.open() as source:
        source.new_page().insert_text((72, 72), text)
        source.save(path)
    with pytest.raises(OCRRequiredError, match="OCR may be required; OCR was not attempted"):
        pdf.extract_pdf(path)


def test_pdf_minimum_text_boundary(tmp_path):
    path = tmp_path / "minimum.pdf"
    with pymupdf.open() as source:
        source.new_page().insert_text((72, 72), "a" * 20)
        source.save(path)
    assert pdf.extract_pdf(path).metadata["native_text_character_count"] == 20


def test_pdf_image_only_does_not_call_ocr(tmp_path, monkeypatch):
    def forbidden_ocr(*args, **kwargs):
        pytest.fail("Native extraction must not invoke OCR")

    monkeypatch.setattr(pymupdf.Page, "get_textpage_ocr", forbidden_ocr)
    image = BytesIO()
    Image.new("RGB", (20, 20), "white").save(image, format="PNG")
    path = tmp_path / "image-only.pdf"
    with pymupdf.open() as source:
        page = source.new_page()
        page.insert_image(pymupdf.Rect(72, 72, 172, 172), stream=image.getvalue())
        source.save(path)
    with pytest.raises(OCRRequiredError):
        extract_document(path)


def test_pdf_reports_textless_pages_without_claiming_complete_extraction(tmp_path):
    path = tmp_path / "partly-empty.pdf"
    with pymupdf.open(FIXTURES / "native.pdf") as source:
        source.new_page()
        source.save(path)
    document = pdf.extract_pdf(path)
    assert document.metadata["pages_without_native_text"] == [3]
    assert document.metadata["page_count"] == 3
    assert len(document.blocks) == 4


def test_password_protected_pdf_is_rejected(tmp_path):
    path = tmp_path / "protected.pdf"
    with pymupdf.open(FIXTURES / "native.pdf") as source:
        source.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256,
                    owner_pw="synthetic-owner", user_pw="synthetic-reader")
    with pytest.raises(ExtractionError, match="Password-protected"):
        pdf.extract_pdf(path)


def test_pdf_failure_after_first_page_never_returns_partial_content(monkeypatch):
    original = pymupdf.Page.get_text

    def fail_on_second_page(page, *args, **kwargs):
        if page.number == 1:
            raise RuntimeError("Synthetic parser failure detail")
        return original(page, *args, **kwargs)

    monkeypatch.setattr(pymupdf.Page, "get_text", fail_on_second_page)
    with pytest.raises(ExtractionError, match="Could not extract PDF") as error:
        pdf.extract_pdf(FIXTURES / "native.pdf")
    assert "Synthetic parser failure detail" not in str(error.value)


def test_docx_preserves_body_and_table_order():
    document = docx.extract_docx(FIXTURES / "native.docx")
    assert [block.text for block in document.blocks] == [
        "Synthetic project overview.", "Item", "State", "Sample widget", "Ready",
        "Synthetic closing paragraph.",
    ]
    assert [document.blocks[index].paragraph_number for index in [0, -1]] == [1, 3]
    assert [block.metadata["body_index"] for block in document.blocks] == [1, 3, 3, 3, 3, 4]
    cells = document.blocks[1:5]
    assert [(block.metadata["row_number"], block.metadata["column_number"])
            for block in cells] == [(1, 1), (1, 2), (2, 1), (2, 2)]
    assert all(block.metadata["table_number"] == 1 for block in cells)
    assert all(block.paragraph_number is None for block in cells)
    assert document.metadata == {"paragraph_count": 3, "table_count": 1}
    assert all(block.extraction_method == "docx_native" for block in document.blocks)
    assert all(block.page_number is block.slide_number is None for block in document.blocks)


@pytest.mark.parametrize("end_cell", [(0, 1), (1, 0)])
def test_docx_merged_cells_are_not_duplicated(tmp_path, end_cell):
    source = DocxDocument()
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(*end_cell)).text = "Merged synthetic cell"
    table.cell(1, 1).text = "Last cell"
    path = tmp_path / "merged.docx"
    source.save(path)
    document = docx.extract_docx(path)
    assert [block.text for block in document.blocks] == ["Merged synthetic cell", "Last cell"]
    assert document.blocks[0].metadata["column_number"] == 1
    assert document.blocks[1].metadata["row_number"] == 2


def test_pptx_preserves_slides_shapes_and_table_cells():
    document = pptx.extract_pptx(FIXTURES / "native.pptx")
    assert [block.text for block in document.blocks] == [
        "Synthetic first slide.", "Second text box.", "Synthetic second slide.",
        "Stage", "State", "Sample task", "Ready",
    ]
    assert [block.slide_number for block in document.blocks] == [1, 1, 2, 2, 2, 2, 2]
    assert [block.metadata["shape_index"] for block in document.blocks] == [1, 2, 1, 2, 2, 2, 2]
    assert document.metadata == {"slide_count": 2}
    assert [(block.metadata["row_number"], block.metadata["column_number"])
            for block in document.blocks[3:]] == [(1, 1), (1, 2), (2, 1), (2, 2)]
    assert all(block.extraction_method == "pptx_native" for block in document.blocks)
    assert all(block.page_number is block.paragraph_number is None for block in document.blocks)


def test_pptx_groups_and_merged_tables_keep_unique_source_paths(tmp_path):
    source = Presentation()
    slide = source.slides.add_slide(source.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(1), Inches(1))  # empty shape retains ordinal
    group = slide.shapes.add_group_shape()
    nested_group = group.shapes.add_group_shape()
    text_shape = nested_group.shapes.add_textbox(0, 0, Inches(1), Inches(1))
    text_shape.text = "Nested synthetic text"
    table_shape = slide.shapes.add_table(2, 2, 0, Inches(2), Inches(3), Inches(2))
    table_shape.table.cell(0, 0).merge(table_shape.table.cell(0, 1))
    table_shape.table.cell(0, 0).text = "Merged synthetic cell"
    path = tmp_path / "grouped.pptx"
    source.save(path)
    document = pptx.extract_pptx(path)
    assert [block.text for block in document.blocks] == ["Nested synthetic text", "Merged synthetic cell"]
    assert document.blocks[0].metadata["shape_path"] == [2, 1, 1]
    assert document.blocks[0].metadata["shape_id"] == text_shape.shape_id
    assert document.blocks[1].metadata["shape_id"] == table_shape.shape_id
    assert document.blocks[1].metadata["shape_path"] == [3]


@pytest.mark.parametrize("file_type", EXTRACTORS)
def test_dispatcher_returns_canonical_document_and_stable_block_ids(file_type):
    path = FIXTURES / f"native.{file_type}"
    document = extract_document(str(path))
    direct = EXTRACTORS[file_type](path)
    assert isinstance(document, Document)
    assert document.filename == path.name
    assert document.file_type == file_type
    assert UUID(document.document_id)
    assert document.document_id != direct.document_id
    assert all(isinstance(block, ContentBlock) and block.text.strip() for block in document.blocks)
    assert document.blocks == direct.blocks
    assert len({block.block_id for block in document.blocks}) == len(document.blocks)
    assert asdict(document)["blocks"][0]["text"] == document.blocks[0].text
    assert "APPROVED" not in repr(document)


@pytest.mark.parametrize("file_type", EXTRACTORS)
def test_dispatcher_accepts_uppercase_extensions(tmp_path, file_type):
    target = tmp_path / f"sample.{file_type.upper()}"
    copyfile(FIXTURES / f"native.{file_type}", target)
    assert extract_document(target).file_type == file_type


@pytest.mark.parametrize("filename", ["sample.txt", "sample.doc", "sample.ppt", "sample", "sample.pdf.exe"])
def test_dispatcher_rejects_unsupported_extensions(filename):
    with pytest.raises(ExtractionError, match="Unsupported file type"):
        extract_document(filename)


@pytest.mark.parametrize("file_type", EXTRACTORS)
def test_missing_file_has_a_project_error(tmp_path, file_type):
    with pytest.raises(ExtractionError, match="file not found"):
        extract_document(tmp_path / f"missing.{file_type}")


@pytest.mark.parametrize("file_type", EXTRACTORS)
def test_corrupt_files_do_not_return_partial_documents_or_expose_content(tmp_path, file_type):
    path = tmp_path / f"corrupt.{file_type}"
    path.write_bytes(b"Synthetic corrupt input marker")
    with pytest.raises(ExtractionError, match="readable and valid") as error:
        extract_document(path)
    assert "Synthetic corrupt input marker" not in str(error.value)
    assert str(path) not in str(error.value)


@pytest.mark.parametrize("file_type", ["docx", "pptx"])
def test_empty_office_files_do_not_claim_success(tmp_path, file_type):
    path = tmp_path / f"empty.{file_type}"
    source = DocxDocument() if file_type == "docx" else Presentation()
    source.save(path)
    with pytest.raises(ExtractionError, match="no extractable"):
        extract_document(path)


@pytest.mark.parametrize("source_type,target_type", [("docx", "pptx"), ("pptx", "docx"), ("docx", "pdf")])
def test_renamed_files_are_rejected_by_the_parser(tmp_path, source_type, target_type):
    path = tmp_path / f"renamed.{target_type}"
    copyfile(FIXTURES / f"native.{source_type}", path)
    with pytest.raises(ExtractionError):
        extract_document(path)
