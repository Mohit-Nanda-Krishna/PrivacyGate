"""Headers/footers, text boxes, speaker notes and embedded-image OCR (synthetic files only)."""

import io
import shutil

import pytest
from docx import Document as DocxDocument
from docx.oxml import parse_xml
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from privacygate.audit import audit_report_to_dict
from privacygate.extraction import OCRRequiredError, embedded, extract_document, ocr, pdf
from privacygate.pipeline import run_pipeline
import pymupdf

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"


def _png(width=800, height=300, text=None) -> io.BytesIO:
    image = Image.new("RGB", (width, height), "white")
    if text:
        ImageDraw.Draw(image).text((20, 20), text, fill="black")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream


def _ocr_data(*lines):
    data = {k: [] for k in ("text", "block_num", "par_num", "line_num", "left", "top", "width", "height")}
    for number, line in enumerate(lines, start=1):
        for index, word in enumerate(line.split()):
            data["text"].append(word)
            data["block_num"].append(1)
            data["par_num"].append(1)
            data["line_num"].append(number)
            data["left"].append(10 + 60 * index)
            data["top"].append(30 * number)
            data["width"].append(50)
            data["height"].append(20)
    return data


@pytest.fixture
def fake_image_ocr(monkeypatch):
    calls = []

    def run(image, label, dpi=None):
        calls.append(label)
        return _ocr_data("Created by Zelda Quark", "Contact z.quark@example.test"), 1

    monkeypatch.setattr(embedded.shutil, "which", lambda command: "tesseract")
    monkeypatch.setattr(embedded, "run_ocr", run)
    return calls


def _docx(tmp_path, *, picture=True, tiny=False, emf=False, filename="embedded.docx"):
    path = tmp_path / filename
    doc = DocxDocument()
    doc.add_paragraph("Body paragraph about widgets.")
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Header owner Ottoline Quarrington"
    section.footer.paragraphs[0].text = "Footer page marker"
    box = parse_xml(
        f'<w:r xmlns:w="{W_NS}" xmlns:v="{V_NS}"><w:pict><v:shape><v:textbox><w:txbxContent>'
        "<w:p><w:r><w:t>Text box contact Hesper Vale</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
    )
    doc.add_paragraph()._p.append(box)
    if picture:
        doc.add_picture(_png(), width=Inches(4))
    if tiny:
        doc.add_picture(_png(120, 60), width=Inches(1))
    doc.save(path)
    if emf:
        # Re-label one media part as EMF to exercise the vector-image warning path.
        import zipfile
        with zipfile.ZipFile(path) as source:
            items = {name: source.read(name) for name in source.namelist()}
        media = [name for name in items if name.startswith("word/media/")][0]
        items[media.rsplit(".", 1)[0] + ".emf"] = items.pop(media)
        for name in list(items):
            if name.endswith(".rels") or name == "[Content_Types].xml":
                items[name] = items[name].replace(media.rsplit("/", 1)[-1].encode(),
                                                  media.rsplit("/", 1)[-1].rsplit(".", 1)[0].encode() + b".emf")
        items["[Content_Types].xml"] = items["[Content_Types].xml"].replace(
            b"</Types>", b'<Default Extension="emf" ContentType="image/x-emf"/></Types>')
        with zipfile.ZipFile(path, "w") as target:
            for name, data in items.items():
                target.writestr(name, data)
    return path


def test_docx_headers_footers_text_boxes_and_images(tmp_path, fake_image_ocr):
    document = extract_document(_docx(tmp_path))
    kinds = {block.metadata.get("kind"): block for block in document.blocks}
    assert kinds["header"].text == "Header owner Ottoline Quarrington"
    assert kinds["footer"].text == "Footer page marker"
    assert kinds["text_box"].text == "Text box contact Hesper Vale"
    image_blocks = [b for b in document.blocks if b.extraction_method == "embedded_image_ocr"]
    assert [b.text for b in image_blocks] == ["Created by Zelda Quark", "Contact z.quark@example.test"]
    assert image_blocks[0].metadata["image_name"].startswith("image")
    assert image_blocks[0].metadata["bbox_pixels"] == [10, 30, 240, 50]  # 4 words, 60 px apart, 50 wide
    assert document.metadata["embedded_images"][0]["status"] == "ocr"
    assert len({b.block_id for b in document.blocks}) == len(document.blocks)


def test_tiny_raster_images_are_enlarged_and_inspected_with_ocr(tmp_path, fake_image_ocr):
    tiny_path = _docx(tmp_path, picture=False, tiny=True, filename="tiny.docx")
    document = extract_document(tiny_path)
    assert len(fake_image_ocr) == 1
    assert document.metadata["embedded_images"][0]["status"] == "ocr"
    assert document.metadata["embedded_images_failed"] == []
    image_blocks = [b for b in document.blocks if b.extraction_method == "embedded_image_ocr"]
    assert len(image_blocks) == 2
    result_tiny = run_pipeline(tiny_path)
    assert result_tiny.validation.status == "APPROVED"


def test_vector_images_are_skipped_with_warnings_and_fail_closed(tmp_path):
    emf_path = _docx(tmp_path, emf=True, filename="vector.docx")
    document = extract_document(emf_path)
    assert document.metadata["embedded_images"][0]["status"] == "skipped_vector"
    assert "vector image (emf)" in document.metadata["extraction_warnings"][0]["reason"]
    assert len(document.metadata["embedded_images_failed"]) == 1
    result_emf = run_pipeline(emf_path)
    assert result_emf.validation.status == "BLOCKED"
    assert "embedded image" in result_emf.validation.reason


def test_tiny_48x48_image_local_ocr_and_gate_behavior(tmp_path, fake_image_ocr):
    path = tmp_path / "tiny_48.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body paragraph with native text.")
    doc.add_picture(_png(48, 48, text="Zelda"), width=Inches(0.5))
    doc.save(path)

    document = extract_document(path)
    assert len(fake_image_ocr) == 1
    assert document.metadata["embedded_images"][0]["status"] == "ocr"
    assert document.metadata["embedded_images_failed"] == []
    result = run_pipeline(path)
    assert result.validation.status == "APPROVED"


def test_tiny_48x48_image_empty_ocr_passes_gate(tmp_path, monkeypatch):
    monkeypatch.setattr(embedded.shutil, "which", lambda command: "tesseract")

    def empty_ocr(image, label, dpi=None):
        return _ocr_data(), 1

    monkeypatch.setattr(embedded, "run_ocr", empty_ocr)

    path = tmp_path / "blank_48.docx"
    doc = DocxDocument()
    doc.add_paragraph("Clean body text with enough native content.")
    doc.add_picture(_png(48, 48), width=Inches(0.5))
    doc.save(path)

    document = extract_document(path)
    assert document.metadata["embedded_images"][0]["status"] == "ocr"
    assert document.metadata["embedded_images"][0]["lines"] == 0
    assert document.metadata["embedded_images_failed"] == []
    result = run_pipeline(path)
    assert result.validation.status == "APPROVED"


def test_tiny_48x48_image_ocr_failure_fails_closed(tmp_path, monkeypatch):
    monkeypatch.setattr(embedded.shutil, "which", lambda command: "tesseract")

    def failing_ocr(image, label, dpi=None):
        raise embedded.ExtractionError(f"OCR failed for {label} (engine error).")

    monkeypatch.setattr(embedded, "run_ocr", failing_ocr)

    path = tmp_path / "fail_48.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body paragraph with native text.")
    doc.add_picture(_png(48, 48), width=Inches(0.5))
    doc.save(path)

    document = extract_document(path)
    assert document.metadata["embedded_images"][0]["status"] == "failed"
    assert len(document.metadata["embedded_images_failed"]) == 1
    result = run_pipeline(path)
    assert result.validation.status == "BLOCKED"
    assert "embedded image" in result.validation.reason


def test_failed_image_ocr_blocks_gate_and_is_audited(tmp_path, monkeypatch):
    monkeypatch.setattr(embedded.shutil, "which", lambda command: "tesseract")

    def broken(image, label, dpi=None):
        raise embedded.ExtractionError(f"OCR failed for {label} (the engine failed). Check Tesseract.")

    monkeypatch.setattr(embedded, "run_ocr", broken)
    result = run_pipeline(_docx(tmp_path), max_passes=2)
    assert result.validation.status == "BLOCKED"
    assert "embedded image" in result.validation.reason
    exported = audit_report_to_dict(result.audit_report)
    assert exported["extraction"]["embedded_images"]["failed"] == 1


def test_missing_tesseract_with_raster_images_fails_closed(tmp_path, monkeypatch):
    monkeypatch.setattr(embedded.shutil, "which", lambda command: None)
    with pytest.raises(OCRRequiredError, match="embedded image"):
        extract_document(_docx(tmp_path))


def test_pptx_notes_and_pictures(tmp_path, fake_image_ocr):
    path = tmp_path / "deck.pptx"
    deck = Presentation()
    for _ in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "Quarterly widgets"
        slide.notes_slide.notes_text_frame.text = "Speaker notes: call Zelda Quark"
        slide.shapes.add_picture(_png(), Inches(1), Inches(2))
    deck.save(path)
    document = extract_document(path)
    notes = [b for b in document.blocks if b.metadata.get("kind") == "notes"]
    assert [(b.slide_number, b.text) for b in notes] == [(1, "Speaker notes: call Zelda Quark"),
                                                         (2, "Speaker notes: call Zelda Quark")]
    images = [b for b in document.blocks if b.extraction_method == "embedded_image_ocr"]
    assert {b.slide_number for b in images} == {1, 2}
    assert len(document.metadata["embedded_images"]) == 1
    assert len(fake_image_ocr) == 1  # the same media file on two slides is OCR'd once


def test_audit_counts_pii_by_location(tmp_path, fake_image_ocr):
    result = run_pipeline(_docx(tmp_path))
    locations = audit_report_to_dict(result.audit_report)["counts_by_location"]
    assert locations.get("image", 0) >= 1  # e-mail found only inside the screenshot


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="Local Tesseract executable unavailable")
def test_real_tesseract_reads_embedded_screenshot(tmp_path):
    path = tmp_path / "real.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body text.")
    image = Image.new("RGB", (900, 300), "white")
    draw = ImageDraw.Draw(image)
    try:
        from PIL import ImageFont
        font = ImageFont.load_default(size=48)
    except TypeError:
        font = None
    draw.text((30, 100), "Synthetic Screenshot Label", fill="black", font=font)
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    doc.add_picture(stream, width=Inches(5))
    doc.save(path)
    document = extract_document(path)
    text = " ".join(b.text for b in document.blocks if b.extraction_method == "embedded_image_ocr")
    assert "Screenshot" in text


def _pdf(tmp_path, *, native_text="Synthetic native paragraph about widgets with sufficient characters.",
         images=1, filename="embedded.pdf", tiny=False):
    path = tmp_path / filename
    with pymupdf.open() as doc:
        page = doc.new_page()
        if native_text:
            page.insert_text((72, 72), native_text)
        for i in range(images):
            stream = _png(120, 60) if tiny else _png(800, 300)
            page.insert_image(pymupdf.Rect(72, 100 + 100 * i, 472, 250 + 100 * i), stream=stream.getvalue())
        doc.save(path)
    return path


def test_pdf_native_page_with_embedded_image_pii(tmp_path, fake_image_ocr):
    path = _pdf(tmp_path)
    document = extract_document(path)
    methods = [b.extraction_method for b in document.blocks]
    assert "pdf_native" in methods
    assert "embedded_image_ocr" in methods

    image_blocks = [b for b in document.blocks if b.extraction_method == "embedded_image_ocr"]
    assert len(image_blocks) == 2
    assert [b.text for b in image_blocks] == ["Created by Zelda Quark", "Contact z.quark@example.test"]
    assert all(b.page_number == 1 for b in image_blocks)
    assert all(b.metadata.get("page_number") == 1 for b in image_blocks)
    assert all(b.metadata.get("kind") == "image" for b in image_blocks)
    assert image_blocks[0].metadata["image_name"].startswith("page_1_img_")
    assert document.metadata["embedded_images"][0]["status"] == "ocr"

    result = run_pipeline(path)
    assert result.validation.status == "APPROVED"
    locations = audit_report_to_dict(result.audit_report)["counts_by_location"]
    assert locations.get("image", 0) >= 1
    sanitized_text = " ".join(b.text for b in result.sanitized_document.blocks)
    assert "Zelda Quark" not in sanitized_text
    assert "z.quark@example.test" not in sanitized_text
    assert "[PERSON" in sanitized_text or "[EMAIL" in sanitized_text


def test_pdf_embedded_image_page_attribution_and_reused_xref(tmp_path, fake_image_ocr):
    path = tmp_path / "reused.pdf"
    stream = _png(800, 300).getvalue()
    with pymupdf.open() as doc:
        p1 = doc.new_page()
        p1.insert_text((72, 72), "First page native text content with enough characters for native threshold.")
        p1.insert_image(pymupdf.Rect(72, 100, 472, 250), stream=stream)
        p1.insert_image(pymupdf.Rect(72, 260, 472, 410), stream=stream)
        p2 = doc.new_page()
        p2.insert_text((72, 72), "Second page native text content with enough characters for native threshold.")
        p2.insert_image(pymupdf.Rect(72, 100, 472, 250), stream=stream)
        doc.save(path)

    document = extract_document(path)
    p1_images = [b for b in document.blocks if b.extraction_method == "embedded_image_ocr" and b.page_number == 1]
    p2_images = [b for b in document.blocks if b.extraction_method == "embedded_image_ocr" and b.page_number == 2]

    # Duplicate on page 1 was deduplicated on page 1
    assert len(p1_images) == 2
    # Page 2 has its own set attributed to page 2
    assert len(p2_images) == 2
    assert all(b.page_number == 1 and b.metadata["page_number"] == 1 for b in p1_images)
    assert all(b.page_number == 2 and b.metadata["page_number"] == 2 for b in p2_images)
    assert all("page_1" in b.block_id for b in p1_images)
    assert all("page_2" in b.block_id for b in p2_images)
    assert len(fake_image_ocr) == 1  # Reused image across pages is OCR'd only once


def test_pdf_embedded_image_ocr_failure_blocks_gate(tmp_path, monkeypatch):
    monkeypatch.setattr(embedded.shutil, "which", lambda command: "tesseract")

    def broken(image, label, dpi=None):
        raise embedded.ExtractionError(f"OCR failed for {label} (engine error). Check Tesseract.")

    monkeypatch.setattr(embedded, "run_ocr", broken)
    path = _pdf(tmp_path)
    result = run_pipeline(path, max_passes=2)
    assert result.validation.status == "BLOCKED"
    assert "embedded image" in result.validation.reason
    exported = audit_report_to_dict(result.audit_report)
    assert exported["extraction"]["embedded_images"]["failed"] == 1


def test_pdf_embedded_image_extraction_failure_blocks_gate(tmp_path):
    path = _pdf(tmp_path)
    original_extract = pymupdf.Document.extract_image

    def failing_extract(self, xref):
        raise RuntimeError("Synthetic corrupted image stream")

    pymupdf.Document.extract_image = failing_extract
    try:
        document = extract_document(path)
        assert len(document.metadata["embedded_images_failed"]) == 1
        assert "page_1_img_" in document.metadata["embedded_images_failed"][0]
        result = run_pipeline(path)
        assert result.validation.status == "BLOCKED"
        assert "embedded image" in result.validation.reason
    finally:
        pymupdf.Document.extract_image = original_extract


def test_pdf_missing_tesseract_with_embedded_images_fails_closed(tmp_path, monkeypatch):
    monkeypatch.setattr(embedded.shutil, "which", lambda command: None)
    path = _pdf(tmp_path)
    with pytest.raises(OCRRequiredError, match="embedded image"):
        extract_document(path)


def test_pdf_tiny_embedded_images_are_inspected_with_ocr(tmp_path, fake_image_ocr):
    path = _pdf(tmp_path, tiny=True)
    document = extract_document(path)
    assert len(fake_image_ocr) == 1
    assert document.metadata["embedded_images"][0]["status"] == "ocr"
    assert len(document.metadata["embedded_images_failed"]) == 0
    result = run_pipeline(path)
    assert result.validation.status == "APPROVED"


def test_pptx_tiny_48x48_image_inspected_and_approves_gate(tmp_path, fake_image_ocr):
    path = tmp_path / "tiny_deck.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Slide with tiny icon"
    slide.shapes.add_picture(_png(48, 48), Inches(1), Inches(1))
    deck.save(path)
    document = extract_document(path)
    assert len(fake_image_ocr) == 1
    assert document.metadata["embedded_images"][0]["status"] == "ocr"
    assert len(document.metadata["embedded_images_failed"]) == 0
    result = run_pipeline(path)
    assert result.validation.status == "APPROVED"


def test_pptx_tiny_image_ocr_failure_blocks_gate(tmp_path, monkeypatch):
    monkeypatch.setattr(embedded.shutil, "which", lambda command: "tesseract")

    def broken(image, label, dpi=None):
        raise embedded.ExtractionError(f"OCR failed for {label} (engine error). Check Tesseract.")

    monkeypatch.setattr(embedded, "run_ocr", broken)
    path = tmp_path / "tiny_deck_fail.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Slide with tiny icon"
    slide.shapes.add_picture(_png(48, 48), Inches(1), Inches(1))
    deck.save(path)
    document = extract_document(path)
    assert document.metadata["embedded_images"][0]["status"] == "failed"
    assert len(document.metadata["embedded_images_failed"]) == 1
    result = run_pipeline(path)
    assert result.validation.status == "BLOCKED"
    assert "embedded image" in result.validation.reason


def test_pdf_mixed_content_coexists_with_scanned_page(tmp_path, fake_image_ocr, monkeypatch):
    path = tmp_path / "mixed_coexist.pdf"
    with pymupdf.open() as doc:
        p1 = doc.new_page()
        p1.insert_image(pymupdf.Rect(0, 0, 400, 600), stream=_png(800, 1200).getvalue())
        p2 = doc.new_page()
        p2.insert_text((72, 72), "Second page native text content with enough characters for native threshold.")
        p2.insert_image(pymupdf.Rect(72, 100, 472, 250), stream=_png(800, 300).getvalue())
        doc.save(path)

    def fake_timed_rec(image, page_number):
        from privacygate.models import ContentBlock
        return [
            ContentBlock(
                block_id=f"page_{page_number}_ocr_block_1",
                text="Scanned page recognized text.",
                page_number=page_number,
                extraction_method="ocr",
                metadata={"block_order": 1},
            )
        ], 1, 0.01, None

    monkeypatch.setattr(pdf, "_timed_recognition", fake_timed_rec)
    monkeypatch.setattr(ocr.shutil, "which", lambda command: "tesseract")
    monkeypatch.setattr(pdf, "render_page", lambda p: Image.new("RGB", (100, 100), "white"))

    document = extract_document(path)
    p1_blocks = [b for b in document.blocks if b.page_number == 1]
    assert len(p1_blocks) == 1
    assert p1_blocks[0].extraction_method == "ocr"
    assert p1_blocks[0].text == "Scanned page recognized text."

    p2_blocks = [b for b in document.blocks if b.page_number == 2]
    p2_methods = {b.extraction_method for b in p2_blocks}
    assert "pdf_native" in p2_methods
    assert "embedded_image_ocr" in p2_methods

    assert document.metadata["ocr_pages"] == [1]
    assert len(document.metadata["embedded_images"]) == 1
    assert "page_2" in document.metadata["embedded_images"][0]["image"]


def test_pdf_reused_tiny_image_deduplicates_failure_reports(tmp_path, monkeypatch):
    monkeypatch.setattr(embedded.shutil, "which", lambda command: "tesseract")

    def broken(image, label, dpi=None):
        raise embedded.ExtractionError(f"OCR failed for {label} (engine error). Check Tesseract.")

    monkeypatch.setattr(embedded, "run_ocr", broken)
    path = tmp_path / "reused_tiny.pdf"
    stream = _png(48, 48).getvalue()
    with pymupdf.open() as doc:
        p1 = doc.new_page()
        p1.insert_text((72, 72), "First page native text content with enough characters for native threshold.")
        p1.insert_image(pymupdf.Rect(72, 100, 192, 160), stream=stream)
        p2 = doc.new_page()
        p2.insert_text((72, 72), "Second page native text content with enough characters for native threshold.")
        p2.insert_image(pymupdf.Rect(72, 100, 192, 160), stream=stream)
        doc.save(path)

    document = extract_document(path)
    assert len(document.metadata["embedded_images_failed"]) == 1
    assert len(document.metadata["embedded_images"]) == 1
    assert document.metadata["embedded_images"][0]["status"] == "failed"
    result = run_pipeline(path)
    assert result.validation.status == "BLOCKED"
    assert "embedded image" in result.validation.reason


def test_pdf_reused_tiny_image_deduplicates_successful_ocr(tmp_path, fake_image_ocr):
    path = tmp_path / "reused_tiny_success.pdf"
    stream = _png(48, 48).getvalue()
    with pymupdf.open() as doc:
        p1 = doc.new_page()
        p1.insert_text((72, 72), "First page native text content with enough characters for native threshold.")
        p1.insert_image(pymupdf.Rect(72, 100, 192, 160), stream=stream)
        p2 = doc.new_page()
        p2.insert_text((72, 72), "Second page native text content with enough characters for native threshold.")
        p2.insert_image(pymupdf.Rect(72, 100, 192, 160), stream=stream)
        doc.save(path)

    document = extract_document(path)
    assert len(fake_image_ocr) == 1
    assert len(document.metadata["embedded_images_failed"]) == 0
    assert len(document.metadata["embedded_images"]) == 1
    assert document.metadata["embedded_images"][0]["status"] == "ocr"
    result = run_pipeline(path)
    assert result.validation.status == "APPROVED"


def test_cadence_pptx_embedded_image_approves_gate():
    from pathlib import Path
    cadence_path = Path("docs/Cadence_Financial_Group_Organizational_Pack.pptx")
    if not cadence_path.exists():
        pytest.skip("Cadence PPTX not present in docs/")
    result = run_pipeline(cadence_path)
    assert result.validation.status == "APPROVED"
    assert len(result.validation.residual_entities) == 0
    assert result.document.metadata["embedded_images_failed"] == []
    report = result.document.metadata["embedded_images"]
    assert len(report) == 1
    assert report[0]["image"] == "image1.png"
    assert report[0]["status"] == "ocr"
