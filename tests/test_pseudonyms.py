"""Stable, document-scoped semantic pseudonyms and privacy-gate boundaries."""

from dataclasses import asdict
from pathlib import Path

from docx import Document as DocxDocument

from privacygate.detection.names import NameRegistry, build_name_registry
from privacygate.models import ContentBlock, Document, PIIEntity, ValidationResult
from privacygate.pipeline import run_pipeline
from privacygate.redaction import PseudonymSession, redact_document, redact_text
from privacygate.validation import validate_privacy


def _document(*texts: str) -> Document:
    return Document(
        "synthetic", "synthetic.docx", "docx",
        blocks=[ContentBlock(f"b{index}", text) for index, text in enumerate(texts)],
    )


def _entity(block: ContentBlock, value: str, entity_type: str = "PERSON") -> PIIEntity:
    start = block.text.index(value)
    return PIIEntity(entity_type, start, start + len(value), 0.9, "test", block.block_id)


def test_repeated_people_and_document_order_with_reverse_splicing() -> None:
    document = _document("Alice Smith met Bob Jones.", "Bob Jones called Alice Smith.")
    entities = [
        _entity(document.blocks[1], "Alice Smith"),
        _entity(document.blocks[0], "Bob Jones"),
        _entity(document.blocks[1], "Bob Jones"),
        _entity(document.blocks[0], "Alice Smith"),
    ]
    session = PseudonymSession()

    sanitized, records = redact_document(document, entities, session=session)

    assert [block.text for block in sanitized.blocks] == [
        "[PERSON_001] met [PERSON_002].",
        "[PERSON_002] called [PERSON_001].",
    ]
    assert [record.placeholder for record in records] == [
        "[PERSON_001]", "[PERSON_002]", "[PERSON_002]", "[PERSON_001]",
    ]
    assert document.blocks[0].text == "Alice Smith met Bob Jones."
    assert "Alice Smith" not in repr(records)
    assert "Bob Jones" not in repr(records)


def test_unique_registry_aliases_reuse_one_token_and_hide_names() -> None:
    document = _document("Miriam T. Achebe", "M. Achebe", "Miriam Achebe")
    registry = build_name_registry(
        document.blocks, [("b0", 0, len(document.blocks[0].text))],
    )
    assert len(registry) == 1
    assert [registry.resolve_person(block.text) for block in document.blocks] == [0, 0, 0]
    session = PseudonymSession(name_registry=registry)

    sanitized, _ = redact_document(
        document, [_entity(block, block.text) for block in document.blocks], session=session,
    )

    assert [block.text for block in sanitized.blocks] == ["[PERSON_001]"] * 3
    for representation in (repr(registry), repr(registry.people), repr(session)):
        assert "Miriam" not in representation
        assert "Achebe" not in representation


def test_different_middle_initials_remain_distinct_and_short_aliases_are_ambiguous() -> None:
    document = _document(
        "Miriam T. Achebe", "Miriam R. Achebe", "Miriam Achebe",
        "M. Achebe", "Miriam T. Achebe",
    )
    registry = build_name_registry(
        document.blocks,
        [(block.block_id, 0, len(block.text)) for block in document.blocks[:2]],
    )
    assert len(registry) == 2
    assert [registry.resolve_person(block.text) for block in document.blocks] == [
        0, 1, None, None, 0,
    ]
    session = PseudonymSession(name_registry=registry)
    entities = [_entity(block, block.text) for block in document.blocks]

    sanitized, _ = redact_document(document, entities, session=session)
    generic, _ = redact_document(document, entities)

    assert [block.text for block in sanitized.blocks] == [
        "[PERSON_001]", "[PERSON_002]", "[PERSON_003]",
        "[PERSON_004]", "[PERSON_001]",
    ]
    assert [block.text for block in generic.blocks] == ["[PERSON]"] * 5


def test_middle_initial_aliases_still_resolve_for_different_first_names() -> None:
    document = _document(
        "Miriam T. Achebe", "Ophelia V. Brandywine", "M. Achebe", "O. Brandywine",
    )
    registry = build_name_registry(
        document.blocks,
        [(block.block_id, 0, len(block.text)) for block in document.blocks[:2]],
    )
    assert [registry.resolve_person(block.text) for block in document.blocks] == [0, 1, 0, 1]

    sanitized, _ = redact_document(
        document, [_entity(block, block.text) for block in document.blocks],
        session=PseudonymSession(name_registry=registry),
    )

    assert [block.text for block in sanitized.blocks] == [
        "[PERSON_001]", "[PERSON_002]", "[PERSON_001]", "[PERSON_002]",
    ]


def test_different_spelled_out_middle_names_are_distinct() -> None:
    document = _document("Miriam Thomas Achebe", "Miriam Travis Achebe")
    registry = build_name_registry(
        document.blocks,
        [(block.block_id, 0, len(block.text)) for block in document.blocks],
        identifier_blocks=[block.block_id for block in document.blocks],
    )
    assert len(registry) == 2
    assert [registry.resolve_person(block.text) for block in document.blocks] == [0, 1]
    assert registry.resolve_person("Miriam T. Achebe") is None


def test_ambiguous_surname_is_not_merged_with_either_person() -> None:
    document = _document(
        "Alice Fairweather", "Bob Fairweather", "Fairweather", "Fairweather",
    )
    registry = build_name_registry(
        document.blocks,
        [("b0", 0, len(document.blocks[0].text)),
         ("b1", 0, len(document.blocks[1].text))],
        identifier_blocks=["b0", "b1"],
    )
    assert len(registry) == 2
    assert registry.resolve_person("Fairweather") is None
    session = PseudonymSession(name_registry=registry)

    sanitized, _ = redact_document(
        document, [_entity(block, block.text) for block in document.blocks], session=session,
    )

    assert [block.text for block in sanitized.blocks] == [
        "[PERSON_001]", "[PERSON_002]", "[PERSON_003]", "[PERSON_003]",
    ]


def test_email_values_type_aliases_and_independent_type_counters() -> None:
    text = "A@Example.COM A@example.com b@example.com EMP-1234 EMP-1234"
    entities = [
        _entity(ContentBlock("b0", text), "EMP-1234", "EMPLOYEE_ID"),
        _entity(ContentBlock("b0", text), "b@example.com", "EMAIL_ADDRESS"),
        _entity(ContentBlock("b0", text), "A@example.com", "EMAIL"),
        _entity(ContentBlock("b0", text), "A@Example.COM", "EMAIL_ADDRESS"),
    ]
    second_id_start = text.rindex("EMP-1234")
    entities.append(PIIEntity("EMPLOYEE_ID", second_id_start, second_id_start + 8, 0.9, "test", "b0"))
    session = PseudonymSession()

    sanitized, _ = redact_text(text, entities, session=session)

    assert sanitized == (
        "[EMAIL_001] [EMAIL_001] [EMAIL_002] [EMPLOYEE_ID_001] [EMPLOYEE_ID_001]"
    )
    assert session.placeholder_for("PHONE_NUMBER", "+1 (555) 010-2030") == "[PHONE_001]"
    assert session.placeholder_for("US_SSN", "123-45-6789") == "[GOVERNMENT_ID_001]"
    assert session.placeholder_for("IBAN_CODE", "GB82 WEST 1234") == "[ACCOUNT_NUMBER_001]"
    assert session.placeholder_for("CUSTOMER_REF", "CUST-1234") == "[CUSTOMER_ID_001]"
    assert session.issued_tokens == frozenset({
        "[EMAIL_001]", "[EMAIL_002]", "[EMPLOYEE_ID_001]", "[PHONE_001]",
        "[GOVERNMENT_ID_001]", "[ACCOUNT_NUMBER_001]", "[CUSTOMER_ID_001]",
    })


def test_new_identity_in_later_pass_gets_next_token_and_legacy_stays_generic() -> None:
    session = PseudonymSession()
    first, _ = redact_text("Alice Smith", [_entity(ContentBlock("b", "Alice Smith"), "Alice Smith")], session)
    second_block = ContentBlock("b", "Bob Jones and Alice Smith")
    second, records = redact_text(
        second_block.text,
        [_entity(second_block, "Bob Jones"), _entity(second_block, "Alice Smith")],
        session,
    )
    generic, _ = redact_text("Alice Smith", [_entity(ContentBlock("b", "Alice Smith"), "Alice Smith")])

    assert first == "[PERSON_001]"
    assert second == "[PERSON_002] and [PERSON_001]"
    assert generic == "[PERSON]"
    assert "Alice Smith" not in repr(session)
    assert "Bob Jones" not in repr(session)
    assert "Alice Smith" not in repr(records)


def test_numbering_extends_beyond_999() -> None:
    session = PseudonymSession()
    for number in range(1, 1001):
        token = session.placeholder_for("PERSON", f"Unique Person {number}")
    assert token == "[PERSON_1000]"
    assert "[PERSON_1000]" in session.issued_tokens


def test_gate_accepts_only_issued_numbered_tokens(monkeypatch) -> None:
    monkeypatch.setattr("privacygate.validation.privacy_gate.detect_pii", lambda document: [])
    document = _document("[PERSON_001] [EMAIL_002] [PERSON] [EMAIL]")

    approved = validate_privacy(
        document, issued_tokens=frozenset({"[PERSON_001]", "[EMAIL_002]"}),
    )
    blocked = validate_privacy(_document("[PERSON_999]"), issued_tokens=frozenset({"[PERSON_001]"}))

    assert approved.status == "APPROVED"
    assert blocked.status == "BLOCKED"
    assert any(entity.detector == "gate:unissued_placeholder" for entity in blocked.residual_entities)
    assert "PERSON_999" not in blocked.reason


def test_gate_ignores_only_detections_contained_inside_an_issued_token(monkeypatch) -> None:
    text = "[PERSON_001]"
    entity = PIIEntity("PERSON", 1, 7, 0.9, "test", "b0")
    monkeypatch.setattr("privacygate.validation.privacy_gate.detect_pii", lambda document: [entity])
    assert validate_privacy(_document(text), issued_tokens=frozenset({text})).status == "APPROVED"

    crossing = PIIEntity("PERSON", 0, len(text) + 1, 0.9, "test", "b0")
    monkeypatch.setattr("privacygate.validation.privacy_gate.detect_pii", lambda document: [crossing])
    assert validate_privacy(_document(text + "X"), issued_tokens=frozenset({text})).status == "BLOCKED"


def test_pipeline_reuses_session_across_residual_passes_without_exporting_it(
    tmp_path: Path, monkeypatch,
) -> None:
    import privacygate.pipeline as pipeline

    path = tmp_path / "synthetic.docx"
    docx = DocxDocument()
    docx.add_paragraph("Alice Smith met Bob Jones.")
    docx.save(path)

    def first_pass_only(document):
        return [_entity(document.blocks[0], "Alice Smith")], NameRegistry()

    seen_issued = []

    def validate(document, **kwargs):
        seen_issued.append(kwargs["issued_tokens"])
        if len(seen_issued) == 1:
            return ValidationResult(
                status="BLOCKED", reason="Residual PII detected.",
                residual_entities=[_entity(document.blocks[0], "Bob Jones")],
            )
        return ValidationResult(status="APPROVED", reason="Clean", residual_entities=[])

    monkeypatch.setattr(pipeline, "detect_pii_with_registry", first_pass_only)
    monkeypatch.setattr(pipeline, "validate_privacy", validate)

    result = run_pipeline(path, max_passes=2)

    assert result.sanitized_document.blocks[0].text == "[PERSON_001] met [PERSON_002]."
    assert seen_issued == [
        frozenset({"[PERSON_001]"}),
        frozenset({"[PERSON_001]", "[PERSON_002]"}),
    ]
    assert result.validation.status == "APPROVED"
    assert result.audit_report.redacted_count == 2
    assert "pseudonyms" not in vars(result)
    assert "tokens_by_identity" not in repr(result.audit_report)
    assert "Alice Smith" not in str(asdict(result.audit_report))
    assert "Bob Jones" not in str(asdict(result.audit_report))
