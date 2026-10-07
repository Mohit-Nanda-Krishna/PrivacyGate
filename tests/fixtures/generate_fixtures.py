"""Regenerate small synthetic native and OCR fixtures with existing dependencies.

Run from the repository root: uv run python tests/fixtures/generate_fixtures.py
These files contain only invented project labels, never real personal data.
"""

from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation
from pptx.util import Inches

FIXTURES = Path(__file__).parent


def generate_ocr_fixtures(directory: Path) -> None:
    """Create deterministic scanned/mixed PDFs; tests generate these in temp dirs.

    Rasterize synthetic text with PyMuPDF's bundled font, then embed only the
    image, without a hidden text layer. No system font or OCR engine is needed.
    """
    with pymupdf.open() as text_document:
        page = text_document.new_page(width=432, height=216)
        page.insert_text((24, 60), "Synthetic scanned report.", fontsize=20)
        page.insert_text((24, 110), "Sample widgets are ready.", fontsize=18)
        png = page.get_pixmap(dpi=150, colorspace=pymupdf.csGRAY).tobytes("png")
    with pymupdf.open() as scanned:
        page = scanned.new_page(width=432, height=216)
        page.insert_image(page.rect, stream=png)
        scanned.save(directory / "scanned.pdf", deflate=True, no_new_id=True)
        with pymupdf.open() as mixed:
            mixed.new_page(width=432, height=216).insert_text(
                (24, 60), "Synthetic native first page.", fontsize=18,
            )
            mixed.insert_pdf(scanned)
            mixed.new_page(width=432, height=216).insert_text(
                (24, 60), "Synthetic native final page.", fontsize=18,
            )
            mixed.save(directory / "mixed.pdf", deflate=True, no_new_id=True)


def main() -> None:
    with pymupdf.open() as pdf:
        first = pdf.new_page()
        # Insert out of visual order to exercise coordinate-sorted extraction.
        first.insert_text((72, 160), "Second block on first page.")
        first.insert_text((72, 72), "Synthetic first page heading.")
        second = pdf.new_page()
        second.insert_text((72, 72), "Synthetic second page heading.")
        second.insert_text((72, 160), "Final block on second page.")
        pdf.set_metadata({"title": "Synthetic native PDF fixture", "author": "PrivacyGate tests"})
        pdf.save(FIXTURES / "native.pdf")

    docx = Document()
    docx.core_properties.author = "PrivacyGate tests"
    docx.core_properties.last_modified_by = "PrivacyGate tests"
    docx.add_paragraph("Synthetic project overview.")
    docx.add_paragraph("")
    table = docx.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("Item", "State"), ("Sample widget", "Ready")]):
        for cell, text in zip(row.cells, values):
            cell.text = text
    docx.add_paragraph("Synthetic closing paragraph.")
    docx.save(FIXTURES / "native.docx")

    pptx = Presentation()
    pptx.core_properties.author = "PrivacyGate tests"
    pptx.core_properties.last_modified_by = "PrivacyGate tests"
    first_slide = pptx.slides.add_slide(pptx.slide_layouts[6])
    first_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Synthetic first slide."
    first_slide.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1)).text = "Second text box."
    second_slide = pptx.slides.add_slide(pptx.slide_layouts[6])
    second_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Synthetic second slide."
    table = second_slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(2)).table
    for row, values in zip(table.rows, [("Stage", "State"), ("Sample task", "Ready")]):
        for cell, text in zip(row.cells, values):
            cell.text = text
    pptx.save(FIXTURES / "native.pptx")
    generate_ocr_fixtures(FIXTURES)


if __name__ == "__main__":
    main()
