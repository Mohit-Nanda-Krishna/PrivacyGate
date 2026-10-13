"""Redacted DOCX/PPTX copies (synthetic files and fake OCR only)."""

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document as WordDocument
from docx.shared import Inches
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches as SlideInches
from streamlit.testing.v1 import AppTest

from privacygate.derivatives import REVIEW_WATERMARK
from privacygate.extraction import embedded
from privacygate.models import ContentBlock, Document, ValidationResult
from privacygate.office_redaction import package_text, redact_office, replacement_pairs
from privacygate.pipeline import run_pipeline

APP_PATH = Path(__file__).parents[1] / "app.py"
NAMES = ("Quillon Ferrant", "Ismay Corvell", "Tobin Read", "Zelda Quarkington")


def _png(width=600, height=200):
    image = Image.new("RGB", (width, height), "white")
    ImageDraw.Draw(image).rectangle((10, 10, 590, 190), outline="black")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


@pytest.fixture
def screenshot_ocr(monkeypatch):
    """The screenshot 'reads' as a ticket created by Tobin Read."""
    def run(image, label, dpi=None):
        words = ["Created", "by", "Tobin", "Read"]
        data = {"text": words, "conf": [95] * 4, "block_num": [1] * 4, "par_num": [1] * 4, "line_num": [1] * 4,
                "left": [20, 120, 160, 240], "top": [40] * 4, "width": [90, 30, 70, 60], "height": [24] * 4}
        return data, 1
    monkeypatch.setattr(embedded.shutil, "which", lambda command: "tesseract")
    monkeypatch.setattr(embedded, "run_ocr", run)


def _docx(path):
    doc = WordDocument()
    doc.core_properties.author = "Quillon Ferrant"
    doc.core_properties.title = "Notes by Quillon Ferrant"
    doc.sections[0].header.paragraphs[0].text = "Owner: Quillon Ferrant"
    paragraph = doc.add_paragraph("The register is owned by ")
    paragraph.add_run("Quillon").bold = True  # one name over two formatted runs
    paragraph.add_run(" Ferrant and reviewed every quarter.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Team", "Members"
    table.cell(1, 0).text = "Risk"
    cell = table.cell(1, 1).paragraphs[0]
    cell.add_run("Ismay Corvell")
    cell.add_run().add_break()  # two names in one cell, separated by a line break
    cell.add_run("Zelda Quarkington")
    doc.add_paragraph("Queue screenshot:")
    doc.add_picture(io.BytesIO(_png()), width=Inches(4))
    doc.inline_shapes[0]._inline.docPr.set("descr", "Ticket created by Tobin Read")
    doc.save(path)
    return path


def _pptx(path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(SlideInches(1), SlideInches(1), SlideInches(6), SlideInches(1)).text_frame.text = \
        "Contact: Ismay Corvell"
    rows = slide.shapes.add_table(2, 2, SlideInches(1), SlideInches(3), SlideInches(6), SlideInches(1)).table
    rows.cell(0, 0).text, rows.cell(0, 1).text = "Owner", "Mobile"
    rows.cell(1, 0).text, rows.cell(1, 1).text = "Quillon Ferrant", "+1 555 0100"
    slide.notes_slide.notes_text_frame.text = "Zelda Quarkington presents this slide."
    deck.save(path)
    return path


def _image_pixels(data: bytes) -> list[tuple[int, int, int]]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        media = [name for name in archive.namelist() if "/media/" in name]
        with Image.open(io.BytesIO(archive.read(media[0]))) as image:
            rgb = image.convert("RGB")
            return [rgb.getpixel((x, y)) for y in range(0, rgb.height, 4) for x in range(0, rgb.width, 4)]


def test_docx_copy_replaces_values_everywhere_and_cleans_properties(tmp_path, screenshot_ocr):
    path = _docx(tmp_path / "memo.docx")
    result = run_pipeline(path)
    copy = redact_office(path, result.document, result.sanitized_document)
    text = package_text(copy.data)
    assert copy.leftover == 0
    for name in NAMES:
        assert name not in text, name
    assert "[PERSON_" in text
    word = WordDocument(io.BytesIO(copy.data))  # still a valid document
    assert word.core_properties.author == "" and word.core_properties.title == ""
    assert any(run.bold for run in word.paragraphs[0].runs)  # formatting of the paragraph survives
    assert set(_image_pixels(copy.data)) == {(0, 0, 0)} and copy.images_masked_whole == 1  # strict default


def test_docx_standard_mode_masks_only_the_redacted_words_in_images(tmp_path, screenshot_ocr):
    path = _docx(tmp_path / "memo.docx")
    result = run_pipeline(path)
    copy = redact_office(path, result.document, result.sanitized_document, mask_all_images=False)
    pixels = _image_pixels(copy.data)
    assert copy.images_masked_whole == 0 and copy.image_words_masked == 2  # "Tobin", "Read"
    assert (0, 0, 0) in set(pixels) and (255, 255, 255) in set(pixels)


def test_pptx_copy_replaces_shapes_tables_and_notes(tmp_path):
    path = _pptx(tmp_path / "deck.pptx")
    result = run_pipeline(path)
    copy = redact_office(path, result.document, result.sanitized_document)
    text = package_text(copy.data)
    assert copy.leftover == 0
    for name in ("Ismay Corvell", "Quillon Ferrant", "Zelda Quarkington", "555 0100"):
        assert name not in text, name
    Presentation(io.BytesIO(copy.data))


def test_blocked_documents_get_a_watermarked_review_copy(tmp_path):
    path = _pptx(tmp_path / "deck.pptx")
    result = run_pipeline(path)
    copy = redact_office(path, result.document, result.sanitized_document, watermark=REVIEW_WATERMARK)
    deck = Presentation(io.BytesIO(copy.data))
    assert all(any(REVIEW_WATERMARK in shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
               for slide in deck.slides)


def test_pairs_follow_the_tokens_even_when_one_chunk_holds_two_names():
    original = Document("d", "x.docx", "docx", blocks=[ContentBlock("b", "Corvell Ismay, Zelda Quarkington")])
    sanitized = Document("d", "x.docx", "docx", blocks=[ContentBlock("b", "[PERSON_001], [PERSON_002]")])
    assert replacement_pairs(original, sanitized) == {"Corvell Ismay": "[PERSON_001]",
                                                      "Zelda Quarkington": "[PERSON_002]"}


def test_only_docx_and_pptx_are_accepted(tmp_path):
    with pytest.raises(ValueError):
        redact_office(tmp_path / "x.pdf", Document("d", "x.pdf", "pdf"), Document("d", "x.pdf", "pdf"))


def test_app_offers_the_redacted_office_download(tmp_path):
    path = _pptx(tmp_path / "deck.pptx")
    result = run_pipeline(path)
    assert result.validation.status == "APPROVED"
    variants = {mode: redact_office(path, result.document, result.sanitized_document,
                                    mask_all_images=mode == "strict") for mode in ("strict", "standard")}
    app = AppTest.from_file(str(APP_PATH))
    app.session_state["pipeline_result"] = result
    app.session_state["redacted_office"] = {result.document.filename: variants}
    app.run(timeout=60)
    assert not app.exception
    assert any("Redacted PowerPoint deck" in str(block.value) for block in app.markdown)
    assert next(box for box in app.checkbox if box.label == "Mask all images").value is True
    assert any("values replaced with the same tokens" in str(caption.value) for caption in app.caption)


def test_app_blocked_office_document_gates_the_review_copy(tmp_path):
    path = _pptx(tmp_path / "deck.pptx")
    result = run_pipeline(path)
    result.validation = ValidationResult(status="BLOCKED", reason="synthetic block")
    copy = redact_office(path, result.document, result.sanitized_document, watermark=REVIEW_WATERMARK)
    app = AppTest.from_file(str(APP_PATH))
    app.session_state["pipeline_result"] = result
    app.session_state["redacted_office"] = {result.document.filename: {"strict": copy, "standard": copy}}
    app.run(timeout=60)
    assert not app.exception
    assert any("redacted copy may still contain personal data" in str(w.value) for w in app.warning)
    assert next(box for box in app.checkbox if box.key == "office_review_copy").value is False
