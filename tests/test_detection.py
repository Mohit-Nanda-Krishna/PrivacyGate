"""Synthetic local PII examples; no real personal records or external services."""

from dataclasses import asdict
import logging
import socket
from unittest.mock import Mock

import pytest
import spacy

import privacygate.detection as detection
from privacygate.detection import DetectionError, detect_pii
from privacygate.detection import presidio_detector
from privacygate.detection.custom_recognizers import detect_custom
from privacygate.detection.regex_detector import detect_regex
from privacygate.extraction import extract_document
from privacygate.models import ContentBlock, Document


def document(text):
    return Document("synthetic", "synthetic.txt", "text", blocks=[ContentBlock("b1", text)])


@pytest.mark.parametrize("text,entity_type,value", [
    ("Alice Johnson lives in London.", "PERSON", "Alice Johnson"),
    ("Alice Johnson lives in London.", "LOCATION", "London"),
    ("Email: alice.johnson@example.com", "EMAIL_ADDRESS", "alice.johnson@example.com"),
    ("Phone: +1 (202) 555-0147", "PHONE_NUMBER", "+1 (202) 555-0147"),
    ("IP address: 192.0.2.10", "IP_ADDRESS", "192.0.2.10"),
    ("Credit card: 4111 1111 1111 1111", "CREDIT_CARD", "4111 1111 1111 1111"),
    # 900 area numbers are unassigned; this is an invented SSN-format test value.
    ("SSN: 900-12-3456", "US_SSN", "900-12-3456"),
    ("IBAN: GB82 WEST 1234 5698 7654 32", "IBAN_CODE", "GB82 WEST 1234 5698 7654 32"),
    ("Employee ID: EMP-29381", "EMPLOYEE_ID", "EMP-29381"),
    ("Employee No: 839201", "EMPLOYEE_ID", "839201"),
    ("Client ID: CLI-29182", "CLIENT_ID", "CLI-29182"),
    ("Customer ID: CUST-83921", "CUSTOMER_ID", "CUST-83921"),
    ("Portfolio ID: PF-29381", "PORTFOLIO_ID", "PF-29381"),
])
def test_hybrid_detects_supported_synthetic_types(text, entity_type, value):
    entities = detect_pii(document(text))
    matches = [entity for entity in entities if entity.entity_type == entity_type]
    assert len(matches) == 1
    entity = matches[0]
    assert text[entity.start:entity.end] == value
    assert entity.block_id == "b1"
    assert 0 < entity.confidence <= 1
    assert entity.risk_level is None
    assert value not in repr(entity)
    assert "text" not in asdict(entity)


@pytest.mark.parametrize("text", [
    "There are 42 widgets and 731 spare parts.",
    "Version 3.11.9 is compatible with version 2.0.1.",
    "The year was 2024 and the next update is in 2026.",
    "Reference AB-839201 and build v29381 are ready.",
    "Microsoft released software in 2024.",
    "The report date is October 7, 2026.",
    "Count 839201. Batch 123456789. Total 19.",
])
def test_harmless_numbers_versions_dates_and_organizations(text):
    assert detect_pii(document(text)) == []


@pytest.mark.parametrize("text,value,source", [
    ("Email: sample@example.com", "sample@example.com", "regex:email"),
    ("Phone: +1 (202) 555-0147", "+1 (202) 555-0147", "regex:labelled_phone"),
])
def test_presidio_and_regex_duplicates_preserve_provenance(text, value, source):
    matches = [entity for entity in detect_pii(document(text)) if text[entity.start:entity.end] == value]
    assert len(matches) == 1
    assert source in matches[0].detector.split("|")
    assert any(name.startswith("presidio:") for name in matches[0].detector.split("|"))


def test_source_locations_block_order_and_unicode_offsets():
    blocks = [
        ContentBlock("z-page", "Résumé: page@example.com", page_number=2,
                     metadata={"block_order": 3, "raw_value": "must not copy"}),
        ContentBlock("a-slide", "Slide: slide@example.com", slide_number=4,
                     metadata={"shape_id": 9, "shape_path": [2, 1], "row_number": 1}),
        ContentBlock("m-paragraph", "Body: body@example.com", paragraph_number=7,
                     metadata={"table_number": 1, "column_number": 2}),
        ContentBlock("ocr-page", "Scan: scan@example.com", page_number=5, extraction_method="ocr",
                     metadata={"source_line_number": 2, "source_block_number": 1}),
    ]
    doc = Document("synthetic", "sample", "pdf", blocks=blocks)
    first = detect_pii(doc)
    assert first == detect_pii(doc)
    assert [entity.block_id for entity in first] == [block.block_id for block in blocks]
    assert [blocks[i].text[entity.start:entity.end] for i, entity in enumerate(first)] == [
        "page@example.com", "slide@example.com", "body@example.com", "scan@example.com",
    ]
    assert first[0].source_location == {"page_number": 2, "block_order": 3}
    assert first[1].source_location == {"slide_number": 4, "shape_id": 9, "shape_path": "2.1", "row_number": 1}
    assert first[2].source_location == {"paragraph_number": 7, "table_number": 1, "column_number": 2}
    assert first[3].source_location == {"page_number": 5, "source_line_number": 2, "source_block_number": 1}


@pytest.mark.parametrize("text,expected", [
    ("Email: test.user+tag@example.internal", "test.user+tag@example.internal"),
    ("<sample@example.com>", "sample@example.com"),
    ("Mobile: +12025550147", "+12025550147"),
])
def test_project_regex_offsets_and_coverage(text, expected):
    results = detect_regex(ContentBlock("b", text))
    assert len(results) == 1
    assert text[results[0].start:results[0].end] == expected


@pytest.mark.parametrize("text", [
    "sample@", "sample@example", "sample..test@example.com", "Version 3.11.9",
    "1234567890", "Phone: 2024-10-07", "xEMP-29381", "Phone: 123",
])
def test_regex_rejects_unstructured_numbers_and_malformed_values(text):
    assert detect_regex(ContentBlock("b", text)) == []


@pytest.mark.parametrize("text,kind,value", [
    ("employee number = 839201", "EMPLOYEE_ID", "839201"),
    ("Employee No.: 839201", "EMPLOYEE_ID", "839201"),
    ("Client ID # C-29182", "CLIENT_ID", "C-29182"),
    ("Customer ID: CUST-83921", "CUSTOMER_ID", "CUST-83921"),
    ("Portfolio ID: AX-29381", "PORTFOLIO_ID", "AX-29381"),
])
def test_enterprise_rules_are_labelled_and_value_only(text, kind, value):
    result, = detect_custom(ContentBlock("b", text))
    assert result.entity_type == kind
    assert text[result.start:result.end] == value


@pytest.mark.parametrize("text", [
    "EMP-29381 CLI-29182 CUST-83921 PF-29381", "Count: 839201", "Employee ID: 2024",
    "Employee ID: EMP-29381suffix", "Employee ID: EMP-29381-extra", "Employee ID: EMP-12345678901",
    "Employee ID:\nEMP-29381", "Employee ID: CLI-29182", "Client ID: unknown",
])
def test_custom_rules_require_context_and_full_value_boundaries(text):
    assert detect_custom(ContentBlock("b", text)) == []


def test_analyzer_has_exact_baseline_and_pinned_model():
    analyzer = presidio_detector.get_analyzer()
    assert set(analyzer.get_supported_entities()) == set(presidio_detector.SUPPORTED_ENTITIES)
    model = analyzer.nlp_engine.nlp["en"]
    assert model.meta["version"] == "3.8.0"
    assert "ner" in model.pipe_names
    assert analyzer.log_decision_process is False


def test_initialization_and_detection_are_offline_and_do_not_log_pii(monkeypatch, caplog):
    def forbidden(*args, **kwargs):
        pytest.fail("Detection must not download or contact external services")

    presidio_detector.get_analyzer.cache_clear()
    monkeypatch.setattr(socket.socket, "connect", forbidden)
    monkeypatch.setattr(spacy.cli, "download", forbidden)
    monkeypatch.setattr(presidio_detector.SpacyNlpEngine, "load", forbidden)
    caplog.set_level(logging.DEBUG)
    entities = detect_pii(document("Alice Johnson lives in London. Email: alice@example.com"))
    assert {entity.entity_type for entity in entities} >= {"PERSON", "LOCATION", "EMAIL_ADDRESS"}
    assert "Alice Johnson" not in caplog.text
    assert "alice@example.com" not in caplog.text


def test_missing_nlp_fails_clearly_without_download(monkeypatch):
    presidio_detector.get_analyzer.cache_clear()
    monkeypatch.setattr(presidio_detector.spacy, "load", Mock(side_effect=OSError("private path")))
    downloader = Mock(side_effect=AssertionError("no downloads"))
    monkeypatch.setattr(spacy.cli, "download", downloader)
    with pytest.raises(DetectionError, match="uv sync --locked") as error:
        detect_pii(document("sample@example.com"))
    assert "private path" not in str(error.value)
    downloader.assert_not_called()
    presidio_detector.get_analyzer.cache_clear()


@pytest.mark.parametrize("name", ["detect_presidio", "detect_regex", "detect_custom"])
def test_detector_failure_never_returns_partial_success(monkeypatch, name):
    monkeypatch.setattr(detection, name, Mock(side_effect=RuntimeError("private text")))
    with pytest.raises(DetectionError, match="no complete result") as error:
        detect_pii(document("sample@example.com"))
    assert "private text" not in str(error.value)


def test_presidio_runtime_failure_has_safe_project_error(monkeypatch):
    engine = Mock(analyze=Mock(side_effect=RuntimeError("private text")))
    monkeypatch.setattr(presidio_detector, "get_analyzer", lambda: engine)
    with pytest.raises(DetectionError, match="Presidio could not analyze") as error:
        presidio_detector.detect_presidio(ContentBlock("b", "sample@example.com"))
    assert "private text" not in str(error.value)


def test_duplicate_block_ids_are_rejected():
    doc = Document("test", "test", "pdf", blocks=[ContentBlock("same", "one"), ContentBlock("same", "two")])
    with pytest.raises(DetectionError, match="unique content block IDs"):
        detect_pii(doc)


def test_empty_document_and_empty_blocks():
    assert detect_pii(Document("empty", "empty", "pdf")) == []
    assert detect_pii(document(" \n\t ")) == []


@pytest.mark.parametrize("file_type", ["pdf", "docx", "pptx"])
def test_extracted_documents_feed_one_detection_interface(tmp_path, file_type):
    text = "Email: sample@example.com"
    path = tmp_path / f"synthetic.{file_type}"
    if file_type == "pdf":
        import pymupdf
        with pymupdf.open() as source:
            source.new_page().insert_text((72, 72), text)
            source.save(path)
    elif file_type == "docx":
        from docx import Document as DocxDocument
        source = DocxDocument()
        source.add_paragraph(text)
        source.save(path)
    else:
        from pptx import Presentation
        from pptx.util import Inches
        source = Presentation()
        slide = source.slides.add_slide(source.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(5), Inches(1)).text = text
        source.save(path)
    doc = extract_document(path)
    entities = detect_pii(doc)
    email, = [entity for entity in entities if entity.entity_type == "EMAIL_ADDRESS"]
    block = next(block for block in doc.blocks if block.block_id == email.block_id)
    assert block.text[email.start:email.end] == "sample@example.com"
