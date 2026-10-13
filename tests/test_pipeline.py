"""Tests for end-to-end PrivacyGate pipeline."""

from pathlib import Path
import pytest
from docx import Document as DocxDocument

from privacygate.extraction.errors import ExtractionError
from privacygate.pipeline import process_document, run_pipeline

FIXTURES = Path(__file__).parent / "fixtures"


def test_pipeline_on_native_fixture_pdf() -> None:
    pdf_path = FIXTURES / "native.pdf"
    result = run_pipeline(pdf_path)

    assert result.document.filename == "native.pdf"
    assert result.validation.status == "APPROVED"
    assert result.audit_report.validation.status == "APPROVED"
    assert len(result.document.blocks) > 0


def test_pipeline_end_to_end_with_pii(tmp_path: Path) -> None:
    # Create a synthetic test DOCX with multiple PII categories
    test_docx = tmp_path / "cadence_report.docx"
    doc = DocxDocument()
    doc.add_paragraph("Employee Jane Doe (Employee ID: EMP-83921) submitted review.")
    doc.add_paragraph("Contact: jane.doe@cadence.com, SSN: 123-45-6789.")
    doc.add_paragraph("Account: ACC-994821.")
    doc.save(test_docx)

    pipeline_result = run_pipeline(test_docx)

    # 1. Extraction checks
    assert pipeline_result.document.file_type == "docx"
    assert len(pipeline_result.document.blocks) == 3

    # 2. Detection checks
    detected_types = {e.entity_type for e in pipeline_result.entities}
    assert "EMPLOYEE_ID" in detected_types
    assert bool(detected_types & {"EMAIL", "EMAIL_ADDRESS"})
    assert bool(detected_types & {"GOVERNMENT_ID", "US_SSN"})
    assert "ACCOUNT_NUMBER" in detected_types

    # 3. Risk classification checks
    for e in pipeline_result.entities:
        assert e.risk_level in ("CRITICAL", "HIGH", "MEDIUM")

    # 4. Redaction checks
    sanitized_text = " ".join(b.text for b in pipeline_result.sanitized_document.blocks)
    assert "[EMPLOYEE_ID_001]" in sanitized_text
    assert "[EMAIL_001]" in sanitized_text
    assert "[GOVERNMENT_ID_001]" in sanitized_text
    assert "[ACCOUNT_NUMBER_001]" in sanitized_text
    assert "123-45-6789" not in sanitized_text
    assert "EMP-83921" not in sanitized_text

    # 5. Secondary Privacy Validation checks
    assert pipeline_result.validation.status == "APPROVED"
    assert len(pipeline_result.validation.residual_entities) == 0

    # 6. Audit Report checks
    audit = pipeline_result.audit_report
    assert audit.detected_count >= 4
    assert audit.redacted_count >= 4
    assert audit.counts_by_risk["CRITICAL"] >= 2  # SSN + Account
    assert audit.counts_by_risk["HIGH"] >= 2      # Employee ID + Email
    assert audit.validation.status == "APPROVED"

    # 7. Check process_document entry point returns matching AuditReport
    direct_report = process_document(test_docx)
    assert direct_report.detected_count == audit.detected_count
    assert direct_report.validation.status == "APPROVED"


def test_pipeline_nonexistent_file_fails_closed(tmp_path: Path) -> None:
    with pytest.raises(ExtractionError):
        run_pipeline(tmp_path / "nonexistent.pdf")


def test_pipeline_multipass_mode(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import privacygate.pipeline
    from privacygate.models import PIIEntity, ValidationResult

    test_docx = tmp_path / "multipass_doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Employee Jane Doe (Employee ID: EMP-83921).")
    doc.save(test_docx)

    # Simulate residual detection on pass 1, followed by clean validation on pass 2
    calls = 0
    orig_validate = privacygate.pipeline.validate_privacy

    def mock_validate(sanitized_doc, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            residual = PIIEntity(
                entity_type="PERSON",
                start=0,
                end=8,
                confidence=0.85,
                detector="mock_detector",
                block_id=sanitized_doc.blocks[0].block_id,
            )
            return ValidationResult(
                status="BLOCKED",
                reason="Residual PII detected: 1 sensitive entity remains.",
                residual_entities=[residual],
            )
        return orig_validate(sanitized_doc, **kwargs)

    monkeypatch.setattr(privacygate.pipeline, "validate_privacy", mock_validate)

    res_single = run_pipeline(test_docx, max_passes=1)
    assert res_single.validation.status == "BLOCKED"
    assert res_single.passes_executed == 1

    calls = 0
    res_multi = run_pipeline(test_docx, max_passes=2)
    assert res_multi.validation.status == "APPROVED"
    assert res_multi.passes_executed == 2


def test_residual_diagnostics_are_separate_from_audit_json(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import privacygate.pipeline as pipeline
    from privacygate.audit import audit_report_to_json
    from privacygate.models import PIIEntity, ValidationResult

    test_docx = tmp_path / "residual_test.docx"
    doc = DocxDocument()
    doc.add_paragraph("Synthetic test document.")
    doc.save(test_docx)

    def mock_validate(sanitized_doc, **kwargs):
        return ValidationResult(
            status="BLOCKED",
            reason="Residual PII detected.",
            residual_entities=[PIIEntity(
                entity_type="GOVERNMENT_ID", start=0, end=8,
                confidence=0.95, detector="gate:id_shape",
                block_id="PRIVATE_BLOCK_ID",
                source_location={
                    "page_number": 7,
                    "extraction_method": "ocr",
                    "filename": "PRIVATE_FILENAME.pdf",
                    "secret": "PRIVATE_SOURCE_VALUE",
                },
            )],
        )

    monkeypatch.setattr(pipeline, "validate_privacy", mock_validate)
    result = pipeline.run_pipeline(test_docx, max_passes=1)
    assert result.validation.status == "BLOCKED"
    assert result.audit_report.validation.status == "BLOCKED"
    assert result.residual_diagnostics == [
        {
            "entity_type": "GOVERNMENT_ID",
            "detectors": ["gate:id_shape"],
            "page": 7,
            "extraction_method": "ocr",
        }
    ]
    exported = audit_report_to_json(result.audit_report)
    assert "residual_diagnostics" not in exported
    assert "PRIVATE_BLOCK_ID" not in exported
    assert "PRIVATE_FILENAME.pdf" not in exported
    assert "PRIVATE_SOURCE_VALUE" not in exported
def test_unresolved_image_ocr_candidate_never_enters_automatic_second_pass(monkeypatch) -> None:
    import privacygate.pipeline as pipeline
    from privacygate.detection.names import NameRegistry
    from privacygate.models import ContentBlock, Document

    block = ContentBlock("image_line", "Quilla Marsh", extraction_method="embedded_image_ocr",
                         metadata={"kind": "image", "image_name": "synthetic.png"})
    document = Document("source", "synthetic.docx", "docx", blocks=[block])
    monkeypatch.setattr(pipeline, "extract_document", lambda path: document)
    monkeypatch.setattr(pipeline, "detect_pii_with_registry", lambda doc: ([], NameRegistry()))
    monkeypatch.setattr("privacygate.validation.privacy_gate.detect_pii", lambda doc: [])

    result = pipeline.run_pipeline("synthetic.docx", max_passes=2)

    assert result.validation.status == "BLOCKED"
    assert result.audit_report.validation.status == "BLOCKED"
    assert "review" in result.validation.reason.lower()
    assert result.passes_executed == 1
    assert result.validation.residual_entities == []
    assert result.redactions == []
    assert result.sanitized_document.blocks[0].text == "Quilla Marsh"
