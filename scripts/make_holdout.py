"""Generate the held-out benchmark: three small synthetic files and their ground truth.

Usage:
    uv run python scripts/make_holdout.py

Writes eval/holdout/docs/ (a DOCX memo with a screenshot, a PPTX deck, a scanned
one-page PDF) and eval/holdout/ground_truth/*.csv. Every name, identifier and
layout here is new: none of it appears in the case-study files, the code or the
tests, so the held-out scores show how the pipeline does on data it was not
tuned on. All values are fictitious; domains use .example, IPs use TEST-NET.
Output is deterministic for a given Pillow / python-docx / python-pptx version.
"""

from __future__ import annotations

import csv
import io
import random
from pathlib import Path

import pymupdf
from docx import Document as WordDocument
from docx.shared import Inches as DocxInches
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "eval" / "holdout"
DOCS = OUT / "docs"
TRUTH = OUT / "ground_truth"
FIELDS = ("file", "page", "location_type", "entity_type", "value", "uncertain", "note")

ORG = "Brightwater Mutual"
DOMAIN = "brightwater-mutual.example"


def font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.load_default(size=size)


def write_truth(name: str, rows: list[tuple[str, str, str, str]]) -> None:
    """rows: (page, location_type, entity_type, value)."""
    TRUTH.mkdir(parents=True, exist_ok=True)
    with open(TRUTH / f"{Path(name).stem}.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(FIELDS)
        for page, location, entity_type, value in rows:
            writer.writerow([name, page, location, entity_type, value, "", ""])


# --- DOCX: memo with an approvals table, a header and a screenshot -----------------

def screenshot() -> bytes:
    """A small 'ticket queue' UI screenshot with an Assignee column."""
    image = Image.new("RGB", (900, 260), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 899, 40), fill=(34, 60, 94))
    draw.text((16, 10), "Service Desk  >  Open tickets", fill="white", font=font(18))
    columns = (16, 150, 420, 640)
    for x, label in zip(columns, ("Ticket", "Summary", "Assignee", "Updated")):
        draw.text((x, 56), label, fill=(90, 90, 90), font=font(15))
    rows = (("SD-1182", "Laptop re-imaging", "Teodoro Bakshi", "03/02/2026"),
            ("SD-1190", "VPN token reset", "Ingrid Haldorsen", "03/04/2026"),
            ("SD-1207", "Shared drive quota", "Noor Elsherbiny", "03/09/2026"))
    for index, row in enumerate(rows):
        y = 92 + index * 50
        draw.line((10, y - 8, 890, y - 8), fill=(225, 225, 225))
        for x, text in zip(columns, row):
            draw.text((x, y), text, fill=(30, 30, 30), font=font(16))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def make_docx() -> None:
    name = "holdout_memo.docx"
    document = WordDocument()
    document.sections[0].header.paragraphs[0].text = f"{ORG} | Facilities change memo | Owner: Callum Ferreira-Wynn"
    document.add_heading("Relocation of the claims operations floor", level=1)
    document.add_paragraph(
        "The claims operations team moves to level 4 on 6 April. Marisol Vantongeren coordinates the move "
        f"and can be reached at marisol.vantongeren@{DOMAIN} or on +1 (503) 555-0849."
    )
    document.add_paragraph(
        "Desk allocations were agreed with the union representative, Kwabena Asante, after the staff "
        "consultation. Personal storage must be cleared by 3 April."
    )
    table = document.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    cells = (("Workstream", "Approver", "Staff ID"),
             ("Seating plan", "Yuki Thornbury", "BWM-48213"),
             ("Network cut-over", "Callum Ferreira-Wynn", "BWM-52914"),
             ("Access badges", "Marisol Vantongeren", "BWM-61432"))
    for r, row in enumerate(cells):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    document.add_paragraph("Current ticket queue for the move:")
    document.add_picture(io.BytesIO(screenshot()), width=DocxInches(6))
    document.add_paragraph(
        "Questions about accessibility adjustments go to the facilities desk; individual needs are "
        "handled confidentially and are not recorded in this memo."
    )
    DOCS.mkdir(parents=True, exist_ok=True)
    document.save(DOCS / name)
    write_truth(name, [
        ("", "body", "PERSON", "Callum Ferreira-Wynn"),
        ("", "prose", "PERSON", "Marisol Vantongeren"),
        ("", "prose", "EMAIL", f"marisol.vantongeren@{DOMAIN}"),
        ("", "prose", "PHONE", "+1 (503) 555-0849"),
        ("", "prose", "PERSON", "Kwabena Asante"),
        ("", "table", "PERSON", "Yuki Thornbury"),
        ("", "table", "EMPLOYEE_ID", "BWM-48213"),
        ("", "table", "EMPLOYEE_ID", "BWM-52914"),
        ("", "table", "EMPLOYEE_ID", "BWM-61432"),
        ("image1.png", "image", "PERSON", "Teodoro Bakshi"),
        ("image1.png", "image", "PERSON", "Ingrid Haldorsen"),
        ("image1.png", "image", "PERSON", "Noor Elsherbiny"),
    ])


# --- PPTX: title slide, contact table, notes ---------------------------------------

def make_pptx() -> None:
    name = "holdout_deck.pptx"
    deck = Presentation()
    title = deck.slides.add_slide(deck.slide_layouts[0])
    title.shapes.title.text = "Quarterly claims review"
    title.placeholders[1].text = "Presented by Ingrid Haldorsen, Head of Claims"
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Escalation contacts"
    table = slide.shapes.add_table(4, 3, Inches(0.6), Inches(1.8), Inches(8.8), Inches(2)).table
    rows = (("Region", "Contact", "Phone"),
            ("North", "Teodoro Bakshi", "+44 161 496 0263"),
            ("South", "Noor Elsherbiny", "+44 117 496 0385"),
            ("West", "Kwabena Asante", "+44 29 2018 0478"))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
            table.cell(r, c).text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(4.3), Inches(8.8), Inches(0.8))
    box.text_frame.text = f"Out of hours: duty manager Yuki Thornbury, yuki.thornbury@{DOMAIN}"
    slide.notes_slide.notes_text_frame.text = (
        "Mention that Marisol Vantongeren replaces the West contact from May."
    )
    DOCS.mkdir(parents=True, exist_ok=True)
    deck.save(DOCS / name)
    write_truth(name, [
        ("1", "body", "PERSON", "Ingrid Haldorsen"),
        ("2", "table", "PERSON", "Teodoro Bakshi"),
        ("2", "table", "PHONE", "+44 161 496 0263"),
        ("2", "table", "PERSON", "Noor Elsherbiny"),
        ("2", "table", "PHONE", "+44 117 496 0385"),
        ("2", "table", "PERSON", "Kwabena Asante"),
        ("2", "table", "PHONE", "+44 29 2018 0478"),
        ("2", "body", "PERSON", "Yuki Thornbury"),
        ("2", "body", "EMAIL", f"yuki.thornbury@{DOMAIN}"),
        ("2", "prose", "PERSON", "Marisol Vantongeren"),
    ])


# --- scanned PDF: one page rendered as an image ----------------------------------

def make_scanned_pdf() -> None:
    name = "holdout_scan.pdf"
    dpi = 200
    page = Image.new("L", (int(8.5 * dpi), int(11 * dpi)), 255)
    draw = ImageDraw.Draw(page)
    x, y = 150, 140
    draw.text((x, y), f"{ORG} - Member incident report", fill=0, font=font(40))
    y += 90
    for line in (
        "Reported by: Callum Ferreira-Wynn (BWM-52914), claims supervisor.",
        "On 11 February a member called about a duplicate premium debit. The member,",
        "Esmond Dunmore, date of birth 22 August 1971, lives at 742 Larkspur Avenue,",
        "Portland, OR 97214 and gave the callback number +1 (971) 555-0591.",
        "His social security number 521-48-3907 was read out on the call and must",
        "be removed from the recording before it is archived.",
    ):
        draw.text((x, y), line, fill=0, font=font(28))
        y += 44
    y += 30
    draw.text((x, y), "Follow-up actions", fill=0, font=font(32))
    y += 60
    columns = (x, x + 520, x + 900)
    for row_index, row in enumerate((("Action", "Owner", "Due"),
                                      ("Refund duplicate debit", "Yuki Thornbury", "14 Feb"),
                                      ("Redact call recording", "Ingrid Haldorsen", "18 Feb"))):
        for column, text in zip(columns, row):
            draw.text((column, y), text, fill=90 if row_index == 0 else 0, font=font(26))
        draw.line((x, y + 40, x + 1350, y + 40), fill=170, width=2)
        y += 56
    y += 50
    # A framed figure: an excerpt of the telephony log.
    box = (x, y, x + 1350, y + 260)
    draw.rectangle(box, outline=200, width=3)
    draw.rectangle((x + 30, y + 30, x + 1320, y + 230), fill=40)
    log = ("2026-02-11 14:02:31 call.inbound from=+1 (971) 555-0591 agent=BWM-48213",
           f"2026-02-11 14:09:57 note.added by=t.bakshi@{DOMAIN} src=198.51.100.207",
           "2026-02-11 14:12:03 recording.flag reason=SSN_SPOKEN")
    for index, line in enumerate(log):
        draw.text((x + 55, y + 60 + index * 52), line, fill=230, font=font(24))
    y += 300
    draw.text((x, y), "Figure 1 - Telephony log excerpt for the call.", fill=60, font=font(22))

    # Scan look: slight blur, noise and rotation.
    rng = random.Random(7)
    noisy = page.filter(ImageFilter.GaussianBlur(0.6)).rotate(0.4, fillcolor=255)
    pixels = noisy.load()
    for _ in range(25000):
        px, py = rng.randrange(noisy.width), rng.randrange(noisy.height)
        pixels[px, py] = max(0, min(255, pixels[px, py] + rng.choice((-60, 60))))
    buffer = io.BytesIO()
    noisy.convert("RGB").save(buffer, format="JPEG", quality=80)
    with pymupdf.open() as pdf:
        pdf_page = pdf.new_page(width=612, height=792)
        pdf_page.insert_image(pdf_page.rect, stream=buffer.getvalue())
        DOCS.mkdir(parents=True, exist_ok=True)
        pdf.save(DOCS / name)
    write_truth(name, [
        ("1", "prose", "PERSON", "Callum Ferreira-Wynn"),
        ("1", "prose", "EMPLOYEE_ID", "BWM-52914"),
        ("1", "prose", "PERSON", "Esmond Dunmore"),
        ("1", "prose", "DATE_OF_BIRTH", "22 August 1971"),
        ("1", "prose", "ADDRESS", "742 Larkspur Avenue, Portland, OR 97214"),
        ("1", "prose", "PHONE", "+1 (971) 555-0591"),
        ("1", "prose", "GOVERNMENT_ID", "521-48-3907"),
        ("1", "table", "PERSON", "Yuki Thornbury"),
        ("1", "table", "PERSON", "Ingrid Haldorsen"),
        ("1", "image", "EMPLOYEE_ID", "BWM-48213"),
        ("1", "image", "EMAIL", f"t.bakshi@{DOMAIN}"),
        ("1", "image", "IP_ADDRESS", "198.51.100.207"),
    ])


def main() -> int:
    make_docx()
    make_pptx()
    make_scanned_pdf()
    print(f"Wrote {sorted(p.name for p in DOCS.iterdir())} and {len(list(TRUTH.glob('*.csv')))} ground truth files")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
